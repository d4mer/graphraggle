"""Worker integration tests for packet-27 (.pptx extracted and submitted as text).

A .pptx must go through ``/documents/text`` with ``file_source`` set to the file
path and the extracted slide/table/notes text as the payload — never through
``/documents/upload``. Extraction failures write a terminal ``failed`` row with
``error_stage="extract"`` and never crash the scan cycle.

Everything is offline: python-pptx builds the fixtures in a temp dir, the
LightRAG client and the state store are fakes.
"""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

# app.worker imports app.config (and app.lightrag_client) at import time, and
# those require these settings to exist. Values are fake; nothing connects.
_REQUIRED_ENV = {
    "RAG_API_KEY": "1234",
    "LIGHTRAG_INTERNAL_API_KEY": "internal",
    "LIGHTRAG_BASE_URL": "http://example:9621",
    "SOURCE_DOCS_DIR": "/tmp/source_docs",
    "UPLOADS_DIR": "/tmp/uploads",
    "STATE_DB_PATH": "/tmp/state.db",
}
for _key, _value in _REQUIRED_ENV.items():
    os.environ.setdefault(_key, _value)

from pptx import Presentation
from pptx.util import Inches

import app.worker as worker


def _run_async(coro):
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    try:
        return loop.run_until_complete(coro)
    finally:
        loop.close()


class _FakeStateStore:
    """In-memory stand-in for the app.state_store functions worker.py calls."""

    def __init__(self) -> None:
        self.docs: dict[str, dict] = {}

    async def get_document_by_path(self, path: str):
        doc = self.docs.get(path)
        return dict(doc) if doc else None

    async def get_documents_by_sha256(self, sha256: str) -> list[dict]:
        return [dict(doc) for doc in self.docs.values() if doc.get("sha256") == sha256]

    async def update_document_state(self, path: str, **kwargs) -> None:
        doc = self.docs.setdefault(path, {"path": path})
        doc.update(kwargs)

    async def upsert_document(self, **kwargs) -> None:
        doc = self.docs.setdefault(kwargs["path"], {"path": kwargs["path"]})
        doc.update(kwargs)

    async def increment_retry(self, path: str) -> int:
        doc = self.docs.get(path)
        if not doc:
            return 0
        doc["retry_count"] = int(doc.get("retry_count") or 0) + 1
        return doc["retry_count"]

    def meta(self, path: str) -> dict:
        raw = self.docs.get(path, {}).get("warnings_json") or "{}"
        block = json.loads(raw).get("worker")
        return block if isinstance(block, dict) else {}


class _FakeClient:
    def __init__(self) -> None:
        self.post_json = AsyncMock(return_value={"track_id": "text-track-1", "status": "accepted"})
        self.post_file = AsyncMock(return_value={"track_id": "upload-track-1", "status": "accepted"})


def _deck_path(directory: Path, name: str = "deck.pptx") -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Q1 revenue"
    slide.placeholders[1].text = "Revenue grew 12%."
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Q1"
    slide.notes_slide.notes_text_frame.text = "Ask for the APAC split."
    path = directory / name
    prs.save(str(path))
    return path


def _garbage_path(directory: Path, name: str = "broken.pptx") -> Path:
    path = directory / name
    path.write_bytes(b"not a pptx, not a zip, just bytes " * 8)
    return path


def _empty_deck_path(directory: Path, name: str = "empty.pptx") -> Path:
    path = directory / name
    Presentation().save(str(path))
    return path


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class _Harness(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.src = self.root / "source_docs"
        self.uploads = self.root / "uploads"
        self.src.mkdir()
        self.uploads.mkdir()
        self.store = _FakeStateStore()
        self.client = _FakeClient()
        self.patcher = patch.multiple(
            worker.settings,
            source_docs_dir=str(self.src),
            uploads_dir=str(self.uploads),
            worker_pptx_extract_enabled=True,
            worker_pptx_max_bytes=104_857_600,
            worker_pptx_max_chars=2_000_000,
        )
        self.patcher.start()
        self.addCleanup(self.patcher.stop)

    def _patches(self):
        return patch.multiple(
            worker,
            get_document_by_path=self.store.get_document_by_path,
            get_documents_by_sha256=self.store.get_documents_by_sha256,
            update_document_state=self.store.update_document_state,
            upsert_document=self.store.upsert_document,
            increment_retry=self.store.increment_retry,
        )

    def _submit(self, path: Path):
        with self._patches(), patch.object(worker, "client", self.client):
            _run_async(worker.submit_file(path, "filesystem", _digest(path)))


# ── Routing ─────────────────────────────────────────────────────────────────

class TestPptxRouting(_Harness):
    def test_uses_pptx_extract_predicate(self):
        self.assertTrue(worker.uses_pptx_extract(self.src / "deck.pptx"))
        self.assertFalse(worker.uses_pptx_extract(self.src / "report.docx"))
        self.assertFalse(worker.uses_pptx_extract(self.src / "notes.txt"))
        with patch.object(worker.settings, "worker_pptx_extract_enabled", False):
            self.assertFalse(worker.uses_pptx_extract(self.src / "deck.pptx"))

    def test_pptx_is_not_treated_as_a_text_extension(self):
        self.assertNotIn(".pptx", worker.TEXT_EXTENSIONS)
        self.assertIn(".pptx", worker.PPTX_EXTENSIONS)

    def test_pptx_goes_to_documents_text_with_the_extracted_payload(self):
        path = _deck_path(self.src)
        self._submit(path)

        self.client.post_file.assert_not_awaited()
        self.client.post_json.assert_awaited_once()
        args = self.client.post_json.await_args.args
        self.assertEqual(args[0], "/documents/text")
        payload = args[1]
        self.assertEqual(payload["file_source"], str(path))
        self.assertIn("## Slide 1: Q1 revenue", payload["text"])
        self.assertIn("Revenue grew 12%.", payload["text"])
        self.assertIn("[Table]", payload["text"])
        self.assertIn("Region | Q1", payload["text"])
        self.assertIn("[Notes]", payload["text"])

        row = self.store.docs[str(path)]
        self.assertEqual(row["status"], "submitted")
        self.assertEqual(row["track_id"], "text-track-1")
        self.assertIsNone(row.get("error_code"))

    def test_flag_off_uses_the_upload_path(self):
        path = _deck_path(self.src)
        with patch.object(worker.settings, "worker_pptx_extract_enabled", False):
            self._submit(path)

        self.client.post_json.assert_not_awaited()
        self.client.post_file.assert_awaited_once()
        self.assertEqual(self.client.post_file.await_args.args[0], "/documents/upload")

    def test_other_binary_extensions_still_upload(self):
        path = self.src / "report.docx"
        path.write_bytes(b"PK\x03\x04 fake docx body")
        self._submit(path)

        self.client.post_file.assert_awaited_once()
        self.assertEqual(self.client.post_file.await_args.args[0], "/documents/upload")
        self.client.post_json.assert_not_awaited()

    def test_scan_cycle_submits_pptx_as_text_and_never_uploads(self):
        _deck_path(self.src)
        with self._patches(), patch.object(worker, "client", self.client):
            _run_async(worker.scan_once())

        self.client.post_file.assert_not_awaited()
        self.client.post_json.assert_awaited_once()
        self.assertEqual(self.client.post_json.await_args.args[0], "/documents/text")

    def test_char_cap_is_applied_before_submitting(self):
        path = _deck_path(self.src)
        with patch.object(worker.settings, "worker_pptx_max_chars", 60):
            self._submit(path)

        submitted = self.client.post_json.await_args.args[1]["text"]
        self.assertIn("[...truncated at 60 characters...]", submitted)


# ── Failure handling ────────────────────────────────────────────────────────

class TestPptxFailures(_Harness):
    def _assert_terminal_extract_failure(self, path: Path, error_code: str) -> dict:
        row = self.store.docs[str(path)]
        self.assertEqual(row["status"], "failed")
        self.assertEqual(row["error_stage"], "extract")
        self.assertEqual(row["error_code"], error_code)
        self.assertIsNone(row.get("track_id"))
        self.assertEqual(self.store.meta(str(path)).get("retryable"), False)
        self.assertEqual(self.store.meta(str(path)).get("failure_class"), "terminal")
        return row

    def test_empty_deck_fails_with_pptx_empty(self):
        path = _empty_deck_path(self.src)
        self._submit(path)

        self._assert_terminal_extract_failure(path, "pptx_empty")
        self.client.post_json.assert_not_awaited()
        self.client.post_file.assert_not_awaited()

    def test_garbage_file_fails_with_pptx_extract_failed(self):
        path = _garbage_path(self.src)
        self._submit(path)

        row = self._assert_terminal_extract_failure(path, "pptx_extract_failed")
        self.assertIn("could not open the file", row["error_detail"])
        self.client.post_json.assert_not_awaited()

    def test_oversize_file_is_refused_before_extraction(self):
        path = _deck_path(self.src)
        with patch.object(worker.settings, "worker_pptx_max_bytes", 10):
            self._submit(path)

        row = self._assert_terminal_extract_failure(path, "pptx_extract_failed")
        self.assertIn("too large", row["error_detail"].lower())
        self.client.post_json.assert_not_awaited()

    def test_missing_python_pptx_reports_pptx_lib_missing(self):
        path = _deck_path(self.src)
        real_import = __import__

        def _blocked_pptx(name, *args, **kwargs):
            if name == "pptx":
                raise ImportError("No module named 'pptx'")
            return real_import(name, *args, **kwargs)

        with patch("builtins.__import__", side_effect=_blocked_pptx):
            self._submit(path)

        self._assert_terminal_extract_failure(path, "pptx_lib_missing")

    def test_scan_cycle_survives_a_failing_pptx(self):
        _garbage_path(self.src)
        with self._patches(), patch.object(worker, "client", self.client):
            _run_async(worker.scan_once())  # must not raise

        row = self.store.docs[str(self.src / "broken.pptx")]
        self.assertEqual(row["status"], "failed")
        self.assertEqual(row["error_code"], "pptx_extract_failed")


if __name__ == "__main__":
    unittest.main()

"""The structure-aware pre-chunker: rules, invariants, and its three call sites.

Covers ``app/structure_chunker.py`` (pure), the worker wiring behind
``STRUCTURE_CHUNKING_ENABLED`` (default off), the experiment-level
``structure_chunking`` key in ``scripts/overnight_loop.py``, and the
``"strategy": "structure"`` arm of ``scripts/chunk_embed_bench.py`` - all the same
function, so a benchmark or loop result describes the shipped path.

All fixtures are synthetic; this suite never reads corpus, probe or LightRAG data.
"""

from __future__ import annotations

import json
import os
import sys
import unittest
from collections import Counter
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_ROOT))
sys.path.insert(0, str(_ROOT / "scripts"))

from app import structure_chunker as sc          # noqa: E402
from app.lightrag_text import (                  # noqa: E402
    is_conservative_marker, pieces_survive, sanitize_like_lightrag)
from app.structure_chunker import (              # noqa: E402
    DELIMITER_CANDIDATES,
    StructureChunkerError,
    chunk_document,
    chunk_token_size_for_payload,
    lightrag_chunk_token_size,
    pick_delimiter,
    render_for_lightrag,
    tokenizer_name,
    token_count,
    _content_words,
    _is_heading,
    coverage_gaps,
    body_words,
)

CONFIGS = _ROOT / "docs" / "ops" / "loop" / "chunk-embed-bench-configs.json"


# ── synthetic fixtures, one per rule ───────────────────────────────────────

TRANSCRIPT = (
    "(Alice Chen) We should ship the reranker this week, the terms are already cleared.\n"
    "(Bob, Acme Legal) I still have concerns about the vendor terms and the cost.\n"
    "(Alice Chen) Legal cleared the terms in July, the memo is in the shared drive.\n"
    "(Bob, Acme Legal) Then we should record the decision before the board meets.\n"
) * 3

HUGE_TURN = TRANSCRIPT + "(Alice Chen) " + ("recorded decision "*90) + "\n"

MARKDOWN_TABLE = ("Table 4 Costs by part\n"
                  "| Part | Cost | Owner |\n|---|---|---|\n"
                  + "\n".join(f"| part {index} | {index * 7} | owner {index} |"
                              for index in range(1, 11)))

TAB_TABLE = ("name\tq1\tq2\tq3\n"
             + "\n".join(f"region {index}\t{index * 3}\t{index * 7}\t{index * 11}"
                         for index in range(1, 12)))

COLUMNAR = ("ColumnA      ColumnB      ColumnC\n"
            + "\n".join(f"value {index}     middle {index}     tail {index}"
                        for index in range(1, 14)))

HEADINGS = """# Procurement policy

## 1. Scope
This policy covers every purchase above the agreed threshold.

### 1.1 Exceptions
Gifts and emergency repairs are handled separately by the owner.

- ask the owner before ordering
- record the amount in the ledger
- keep the receipt attached to the claim

## 2. Approval
Two signatures are required above the cap, one of them the budget holder.
"""

NESTED_LISTS = ("Lead-in line that introduces what the list below means.\n\n"
                "- first item with a few extra words so it is not trivial\n"
                "  - nested one.a under the first item\n"
                "  - nested one.b under the first item\n"
                "- second item with its own tail of words\n\n"
                "- third item after the blank line\n")

FURNITURE = "\n".join(
    ["Acme Corp confidential"] * 5
    # a running footer that is identical on every page is the recognisable shape;
    # "Page 1", "Page 2", ... are not repeats and must never be dropped
    + ["Page 3"] * 5
    + ["Section Alpha", "Alpha alpha alpha alpha alpha.", "Section Beta",
       "Beta beta beta beta beta beta."] * 3
)

TIMESTAMPISH = """1.2.3 Numbers

0:12 12 2026-01-02

Reported totals were 12 and 20 for the quarter.
"""

CRLF_UNICODE = ("Header line\r\n\r\nCafé naïve — “quoted” and é and ü.\r\n"
                + "Deux mots encore. " * 12 + "\r\nFin.\r\n")

PLAIN = ("It was a Tuesday when the decision was finally taken, after weeks of argument "
         "about cost and about the vendor's terms. ") * 6


# ── the rules, one test each ───────────────────────────────────────────────

class TestRules(unittest.TestCase):
    def test_transcript_is_cut_at_turn_boundaries_and_detected(self):
        plan = chunk_document(TRANSCRIPT, filename="Standup_2026-10-09.txt",
                              budget_tokens=80, header_tokens=25)
        self.assertEqual(plan.kind_detected, "transcript")
        self.assertGreater(len(plan.pieces), 1)
        for piece in plan.pieces:
            self.assertEqual(piece.kind, "turn")
            lines = [line for line in piece.body.splitlines() if line.strip()]
            self.assertTrue(lines, "empty piece")
            self.assertTrue(lines[0].startswith("("), "a piece must not start mid-turn")
            self.assertIn("[Meeting: Standup | Date: 2026-10-09", piece.header)
        self.assertEqual({piece.part_total for piece in plan.pieces}, {len(plan.pieces)})

    def test_one_huge_turn_is_split_without_losing_a_word(self):
        plan = chunk_document(HUGE_TURN, filename="Standup_2026-10-09.txt",
                              budget_tokens=120, header_tokens=20)
        self.assertEqual(plan.kind_detected, "transcript")
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, 120)
        self.assertEqual(plan.stats.get("oversized_turns_split"), 1)
        # the giant turn's words are all still there
        bodies = " ".join(piece.body for piece in plan.pieces)
        self.assertEqual(bodies.count("recorded decision"), 90)

    def test_markdown_table_splits_by_row_with_the_header_row_repeated(self):
        plan = chunk_document(MARKDOWN_TABLE, filename="Costs_2026-04-11.md",
                              budget_tokens=90, header_tokens=20)
        self.assertEqual(plan.kind_detected, "table")
        self.assertGreater(len(plan.pieces), 1)
        for index, piece in enumerate(plan.pieces):
            lines = piece.body.splitlines()
            if index:
                self.assertEqual(lines[0], "| Part | Cost | Owner |",
                                 "every piece of a split table carries the header row")
                self.assertEqual(lines[1], "|---|---|---|", "and its separator rule")
        self.assertGreaterEqual(plan.stats.get("repeated_table_header_rows", 0), 2)
        # the caption stays with the FIRST piece only
        self.assertIn("Table 4 Costs by part", plan.pieces[0].body)
        self.assertNotIn("Table 4 Costs by part",
                         " ".join(piece.body for piece in plan.pieces[1:]))

    def test_tab_separated_table_is_recognised_as_a_table(self):
        plan = chunk_document(TAB_TABLE, filename="Regions_2026.tsv", budget_tokens=70,
                              header_tokens=16)
        self.assertEqual(plan.kind_detected, "table")
        self.assertEqual(plan.stats.get("tables"), 1)
        for piece in plan.pieces[1:]:
            self.assertEqual(piece.body.splitlines()[0], "name\tq1\tq2\tq3")

    def test_columnar_two_space_runs_are_a_table(self):
        plan = chunk_document(COLUMNAR, filename="Wide.txt", budget_tokens=70,
                              header_tokens=16)
        self.assertEqual(plan.kind_detected, "table")
        self.assertTrue(all(piece.kind == "table" for piece in plan.pieces))

    def test_a_row_wider_than_the_budget_splits_at_cell_boundaries_repeating_the_key(self):
        wide = ("| Item | Notes |\n|---|---|\n"
                "| key-cell | " + " | ".join(f"note {index}" for index in range(40)) + " |\n")
        plan = chunk_document(wide, filename="Wide_2026.md", budget_tokens=60,
                              header_tokens=12)
        bodies = [piece.body for piece in plan.pieces]
        self.assertGreater(len(bodies), 1)
        for body in bodies:
            self.assertIn("key-cell", body, "the row key is repeated on every piece")
        self.assertGreaterEqual(plan.stats.get("repeated_table_row_keys", 0), 1)
        self.assertNotIn("cell-over-budget", plan.warnings,
                         "cell splitting avoided splitting inside a cell")

    def test_a_single_huge_cell_is_split_by_words_and_says_so(self):
        blob = ("| Item | Notes |\n|---|---|\n"
                "| key-cell | " + " ".join(f"word{index}" for index in range(60)) + " |\n")
        plan = chunk_document(blob, filename="Blob_2026.md", budget_tokens=60,
                              header_tokens=12)
        self.assertGreater(len(plan.pieces), 1)
        self.assertEqual(plan.stats.get("table_cells_split"), 1)
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, 60)

    def test_headings_start_pieces_and_feed_the_section_path(self):
        plan = chunk_document(HEADINGS, filename="Procurement_2026-03-15.md",
                              budget_tokens=70, header_tokens=30)
        self.assertEqual(plan.kind_detected, "structured")
        paths = [piece.section_path for piece in plan.pieces]
        self.assertIn(("Procurement policy", "Scope"), paths)
        self.assertIn(("Procurement policy", "Scope", "Exceptions"), paths)
        self.assertIn(("Procurement policy", "Approval"), paths)
        headed = [piece for piece in plan.pieces if piece.kind == "section"]
        self.assertTrue(headed)
        self.assertTrue(any("Section: Procurement policy / Scope" in (piece.header or "")
                            for piece in headed))

    def test_a_list_travels_with_its_lead_in_when_it_fits(self):
        plan = chunk_document(NESTED_LISTS, filename="Lists.txt", budget_tokens=200,
                              header_tokens=20)
        self.assertEqual(len(plan.pieces), 1)
        self.assertIn("Lead-in line", plan.pieces[0].body)
        self.assertIn("nested one.a", plan.pieces[0].body)
        tight = chunk_document(NESTED_LISTS, filename="Lists.txt", budget_tokens=45,
                               header_tokens=10)
        self.assertGreater(len(tight.pieces), 1)
        for piece in tight.pieces:
            self.assertLessEqual(piece.token_estimate, 45)

    def test_running_headers_and_page_markers_are_dropped_and_counted(self):
        plan = chunk_document(FURNITURE, filename="Deck.txt", budget_tokens=90,
                              header_tokens=20)
        self.assertEqual(plan.stats.get("dropped_running_lines"), 16)
        self.assertEqual(plan.stats.get("dropped_page_markers"), 5)
        joined = "\n".join(piece.body for piece in plan.pieces)
        self.assertNotIn("Acme Corp confidential", joined)
        self.assertNotIn("Page 3", joined)
        # ... and the content that only LOOKED repeatable is intact
        self.assertEqual(joined.count("Alpha alpha alpha alpha alpha."), 3)
        self.assertEqual(joined.count("Beta beta beta beta beta beta."), 3)

    def test_repeated_content_lines_are_never_dropped(self):
        repeated = "(Alice Chen) Got it.\n" * 6 + "(Bob, Acme Legal) Understood here too.\n" * 6
        plan = chunk_document(repeated, filename="Standup.txt", budget_tokens=120,
                              header_tokens=20)
        joined = "\n".join(piece.body for piece in plan.pieces)
        self.assertEqual(joined.count("Got it."), 6)
        self.assertEqual(joined.count("Understood here too."), 6)
        self.assertEqual(plan.stats.get("dropped_running_lines", 0), 0)

    def test_numbers_that_look_like_timestamps_are_not_headings_or_furniture(self):
        plan = chunk_document(TIMESTAMPISH, filename="Numbers_2026-01-02.txt",
                              budget_tokens=120, header_tokens=20)
        joined = "\n".join(piece.body for piece in plan.pieces)
        self.assertIn("0:12 12 2026-01-02", joined)
        self.assertIn("Reported totals were 12 and 20", joined)
        self.assertEqual(plan.stats.get("dropped_running_lines", 0), 0)

    def test_crlf_and_unicode_survive_intact(self):
        plan = chunk_document(CRLF_UNICODE, filename="Notes_2026-02-02.txt",
                              budget_tokens=60, header_tokens=20)
        joined = "\n".join(piece.body for piece in plan.pieces)
        self.assertIn("Café naïve — “quoted” and é and ü.", joined)
        self.assertNotIn("\r", joined)
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, 60)

    def test_plain_text_falls_back_to_paragraphs_sentences_words(self):
        plan = chunk_document(PLAIN, filename="notes.txt", budget_tokens=50,
                              header_tokens=15)
        self.assertEqual(plan.kind_detected, "plain")
        self.assertGreater(len(plan.pieces), 1)
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, 50)
        self.assertEqual(" ".join(piece.body for piece in plan.pieces).count(
            "decision was finally taken"), 6)

    def test_coverage_accounts_for_every_word_or_reports_the_gap(self):
        from app.structure_chunker import ChunkPlan, Piece, coverage_gaps
        source = "alpha beta gamma delta epsilon zeta\n"
        full = chunk_document(source, filename="words.txt", budget_tokens=12, header_tokens=4)
        whole = coverage_gaps(source, full)
        self.assertEqual(whole["unaccounted_word_tokens"], 0)
        self.assertEqual(whole["covered_pct"], 100.0)
        # A plan that lost words must say how many, in words, and not name them.
        leaky = ChunkPlan(pieces=[Piece(body="alpha beta", header="", kind="plain",
                                        section_path=())], kind_detected="plain")
        gaps = coverage_gaps(source, leaky)
        self.assertEqual(gaps["missing_word_tokens"], 4)
        self.assertEqual(gaps["unaccounted_word_tokens"], 4)
        self.assertLess(gaps["covered_pct"], 100.0)
        blob = json.dumps(gaps)
        for leaked in ("gamma", "delta", "epsilon", "zeta"):
            self.assertNotIn(leaked, blob, "the coverage report carries counts, never text")

    def test_no_input_shape_can_stall_the_scan(self):
        """A line none of the branches claims must still be consumed, not looped on.

        A hang here is an overnight ingest that never finishes, which is worse than a
        wrong chunk: the driver's only defence against an ingest that never returns is
        its wall-clock ceiling, and it would spend the whole night there. (Found by
        sabotage: disable the heading branch and the scan never advances on "# Heading".)
        """
        odd = ["# Heading", "* item", "  ", "\t", "|", "|---|", "1.2.3.4.5.6.7.8 x",
               "---", "   \t  x", "\f", "-"]
        text = "\n".join(odd * 3)
        plan = chunk_document(text, filename="odd.txt", budget_tokens=40, header_tokens=12)
        self.assertTrue(plan.pieces)
        joined = "\n".join(piece.body for piece in plan.pieces)
        # Structure survives repetition: a heading or a list item that appears three
        # times is the document repeating itself, not a page header.
        self.assertIn("# Heading", joined)
        self.assertIn("* item", joined)
        # The only thing allowed to disappear is the furniture-shaped line, and it is
        # counted rather than silently removed.
        kept = [word for line in text.splitlines() if line.strip()
                for word in _content_words(line)]
        out = [word for piece in plan.pieces for word in _content_words(piece.body)]
        self.assertEqual(Counter(kept) - Counter(out),
                         Counter({"1.2.3.4.5.6.7.8": 3, "x": 3}),
                         "only the furniture-shaped line may go")
        self.assertEqual(plan.stats.get("dropped_running_lines"), 3)

    def test_a_forced_kind_that_does_not_fit_says_so_and_still_returns_a_plan(self):
        plan = chunk_document(PLAIN, filename="notes.txt", kind="transcript",
                              budget_tokens=50)
        self.assertIn("transcript-not-detected", plan.warnings)
        self.assertEqual(plan.kind_detected, "plain")
        self.assertTrue(plan.pieces)

    def test_kind_can_be_forced_off_for_speed(self):
        plan = chunk_document(TRANSCRIPT, filename="Standup.txt", kind="plain",
                              budget_tokens=100)
        self.assertNotEqual(plan.kind_detected, "transcript")

    def test_empty_input_returns_an_empty_plan_rather_than_raising(self):
        plan = chunk_document("   \n\n", filename="blank.txt")
        self.assertTrue(plan.empty)
        self.assertIn("empty-input", plan.warnings)

    def test_an_unsplittable_word_is_reported_not_silently_truncated(self):
        monster = "short. " + ("x" * 400) + " tail.\n" * 2
        plan = chunk_document(monster, filename="blob.txt", budget_tokens=60,
                              header_tokens=10)
        oversized = [piece for piece in plan.pieces if piece.token_estimate > 60]
        self.assertTrue(oversized, "the unsplittable word must survive, oversize")
        for piece in oversized:
            self.assertEqual(len(piece.body.split()), 1,
                             "only a single unsplittable word may exceed the budget")


# ── the invariants, across the whole fixture set ──────────────────────────

FIXTURES = {
    "transcript": TRANSCRIPT, "huge turn": HUGE_TURN, "markdown table": MARKDOWN_TABLE,
    "tab table": TAB_TABLE, "columnar": COLUMNAR, "headings": HEADINGS,
    "nested lists": NESTED_LISTS, "furniture": FURNITURE, "timestamps": TIMESTAMPISH,
    "crlf unicode": CRLF_UNICODE, "plain": PLAIN,
}


class TestNaturalUnitsNoCap(unittest.TestCase):
    """The operator's shape: no chunk-size cap, the embedder window is the only ceiling.

    Hermes' real-document run is the reason these exist: a 350-token cap turned one PDF
    into 596 pieces with a median of 111 tokens, because every heading and table boundary
    forced a cut. A section that fits the window is now handed over whole, and a page of
    small sections welds together instead of shattering.
    """

    @staticmethod
    def _sentences(count: int, *, topic: str = "budget") -> str:
        words = ("allocation freight region ledger quarter margin carrier volume "
                 "reserve tariff corridor depot intake output").split()
        return " ".join(
            f"{topic.title()} line {index} records the {words[index % len(words)]} figure "
            f"for quarter {index % 4 + 1} of {2020 + index % 6}."
            for index in range(count))

    def test_a_section_that_fits_the_window_is_never_cut(self):
        text = "# Quarterly review\n\n" + self._sentences(380) + "\n"
        plan = chunk_document(text, filename="Review_2026-04.txt", max_piece_tokens=7500,
                              min_piece_tokens=150)
        self.assertEqual(len(plan.pieces), 1,
                         "7000 tokens of one section is ONE piece: nothing may cut it")
        self.assertGreaterEqual(plan.stats["tokens_max"], 4000,
                                "the fixture has to be big enough to make the point")
        self.assertLessEqual(plan.stats["tokens_max"], 7500)
        self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)

    def test_a_section_bigger_than_the_window_splits_at_its_sub_headings(self):
        parts = ["# Annual annex"]
        for section in range(1, 5):
            parts.append(f"## {section}. {['Freight', 'Intake', 'Tariffs', 'Depots'][section - 1]}")
            parts.append(self._sentences(170, topic=["freight", "intake", "tariffs",
                                                     "depots"][section - 1]))
        text = "\n\n".join(parts) + "\n"
        plan = chunk_document(text, filename="Annex_2026-04.txt", max_piece_tokens=7500,
                              min_piece_tokens=150)
        self.assertGreater(len(plan.pieces), 1, "the window is what forced the cut")
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, 7500)
        bodies = [piece.body for piece in plan.pieces]
        self.assertTrue(any(body.startswith("# Annual annex") for body in bodies),
                        "the parent heading is not lost when its section splits")
        for heading in ("## 2. Intake", "## 3. Tariffs", "## 4. Depots"):
            self.assertTrue(any(body.startswith(heading) for body in bodies),
                            f"{heading} starts a piece: the split is at the sub-heading, "
                            "not at a token count")

    def test_a_heading_never_stands_alone(self):
        text = ("# Procurement\n\n## Scope\n\n"
                + self._sentences(24, topic="scope")
                + "\n\n## Approval\n\n" + self._sentences(24, topic="approval") + "\n")
        plan = chunk_document(text, filename="Policy.txt", max_piece_tokens=700,
                              min_piece_tokens=150)
        for piece in plan.pieces:
            self.assertNotEqual(len(piece.body.strip().splitlines()), 1,
                                f"a piece is a lone heading: {piece.body!r}")
            self.assertFalse(all(_is_heading(line) for line in piece.body.splitlines()),
                             "a heading without its content is retrievable by nothing")

    def test_a_run_of_small_sections_does_not_fragment(self):
        parts = []
        for section in range(1, 61):
            parts.append(f"{section}. {['Freight', 'Intake', 'Tariffs', 'Depots', 'Wages'][section % 5]}"
                         f" note {section}")
            parts.append(self._sentences(7, topic=str(section)))
        text = "\n\n".join(parts) + "\n"
        default = chunk_document(text, filename="Notes.txt", max_piece_tokens=7500,
                                 min_piece_tokens=150)
        total = sum(piece.token_estimate for piece in default.pieces)
        self.assertGreater(total, 4000, "fixture too small to model a real document")
        self.assertEqual(default.stats["merge_until"], "floor",
                         "the lead's call: a merge chain stops at the floor, so a piece "
                         "stays one natural unit instead of a mixed-topic bundle")
        self.assertLess(len(default.pieces), len(chunk_document(
            text, filename="Notes.txt", max_piece_tokens=350, min_piece_tokens=0).pieces),
            "the same document under Hermes' 350-token cap shatters; the floor does not")
        for piece in default.pieces[:-1]:
            self.assertGreaterEqual(piece.token_estimate, 150,
                                    "no fragment is left below the floor")
        chained = chunk_document(text, filename="Notes.txt", max_piece_tokens=7500,
                                 min_piece_tokens=150, until_merge="window")
        self.assertLess(len(chained.pieces), len(default.pieces),
                        "window mode remains an opt-in that runs the chain to the window")
        for piece in chained.pieces[:-1]:
            self.assertGreaterEqual(piece.token_estimate, 150)
        for plan in (default, chained):
            self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)

    def test_a_parent_section_is_one_unit_when_its_whole_extent_fits(self):
        parts = ["# Freight annex"]
        for child in range(1, 6):
            parts.append(f"## {child}. {'Depot intake tariffs wages routes'.split()[child - 1]}")
            parts.append(self._sentences(9, topic=str(child)))
        text = "\n\n".join(parts) + "\n"
        plan = chunk_document(text, filename="Annex.txt", max_piece_tokens=7500,
                              min_piece_tokens=150)
        self.assertEqual(len(plan.pieces), 1,
                         "a parent that fits the window is its natural extent: the "
                         "siblings under it do not become five pieces")
        self.assertEqual(plan.stats["natural_units"], 1)

    def test_a_table_that_fits_the_window_is_not_split_and_its_header_is_not_repeated(self):
        rows = ["| Region | Q1 | Q2 | Q3 |", "|---|---|---|---|"] + [
            f"| depot {index:02d} | {index * 3} | {index * 5} | {index * 7} |"
            for index in range(1, 45)]
        text = "# Order book\n\n" + "\n".join(rows) + "\n\n" + self._sentences(10) + "\n"
        plan = chunk_document(text, filename="Orderbook_2026-04.txt", max_piece_tokens=7500,
                              min_piece_tokens=150)
        self.assertEqual(len(plan.pieces), 1, "44 rows under the window is one table")
        self.assertEqual(plan.stats.get("tables_split_by_row", 0), 0,
                         "no row splitting, so no invented repeated header")
        self.assertEqual(plan.pieces[0].body.count("| Region | Q1 | Q2 | Q3 |"), 1)

    def test_the_window_still_splits_a_table_that_cannot_fit(self):
        rows = ["| Region | Q1 | Q2 | Q3 |", "|---|---|---|---|"] + [
            f"| depot {index:03d} | {index * 3} | {index * 5} | {index * 7} |"
            for index in range(1, 1201)]
        text = "# Order book\n\n" + "\n".join(rows) + "\n"
        plan = chunk_document(text, filename="Orderbook_Large_2026-04.txt",
                              max_piece_tokens=7500, min_piece_tokens=150)
        self.assertGreater(len(plan.pieces), 1)
        self.assertGreaterEqual(plan.stats.get("tables_split_by_row", 0), 1)
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, 7500)
            self.assertIn("| Region | Q1 | Q2 | Q3 |", piece.body,
                          "a piece of a split table never strands a number from its columns")

    def test_the_ceiling_and_the_floor_are_reported_not_guessed(self):
        plan = chunk_document(self._sentences(60) + "\n", filename="Plain.txt",
                              max_piece_tokens=900, min_piece_tokens=200)
        self.assertEqual(plan.stats["max_piece_tokens"], 900)
        self.assertEqual(plan.stats["min_piece_tokens"], 200)
        self.assertEqual(plan.stats["merge_until"], "floor", "the production default")


class TestInvariants(unittest.TestCase):
    """The contract, checked over every fixture at three budgets.

    Bodies, minus inserted headers and the repeats the rules force, account for every
    content word of the input; nothing is over budget unless it is one unsplittable
    word; no piece is empty; the delimiter appears nowhere inside a body; and the
    rendered text splits back into exactly the planned pieces.
    """

    def check(self, name, text, budget, headers=True):
        plan = chunk_document(text, filename="Vendor_Review_2026-04-11.txt",
                              budget_tokens=budget, header_tokens=25, headers=headers)
        self.assertTrue(plan.pieces, f"{name} b{budget}: no pieces")
        kept = [word for line in text.splitlines() if line.strip()
                for word in _content_words(line)]
        dropped = plan.stats.get("dropped_running_words", 0)
        repeats = plan.stats.get("repeated_words", 0)
        out = [word for piece in plan.pieces for word in _content_words(piece.body)]
        missing = Counter(kept) - Counter(out)
        self.assertEqual(sum(missing.values()), dropped,
                         f"{name} b{budget}: lost words {dict(missing)}")
        self.assertEqual(len(out), len(kept) - dropped + repeats,
                         f"{name} b{budget}: accounting out={len(out)} kept={len(kept)} "
                         f"dropped={dropped} repeats={repeats}")
        for piece in plan.pieces:
            self.assertTrue(piece.body.strip(), f"{name} b{budget}: empty piece")
            if piece.token_estimate > budget:
                self.assertEqual(len(piece.body.split()), 1,
                                 f"{name} b{budget}: oversize multi-word piece "
                                 f"{piece.body[:60]!r}")
        rendered, params = render_for_lightrag(plan)
        delimiter = params["params"]["split_by_character"]
        self.assertNotIn(delimiter, "".join(piece.body for piece in plan.pieces),
                         f"{name} b{budget}: delimiter inside a body")
        parts = [part for part in rendered.split(delimiter) if part]
        self.assertEqual(len(parts), len(plan.pieces),
                         f"{name} b{budget}: render split into {len(parts)}")
        for piece, part in zip(plan.pieces, parts):
            self.assertEqual(part, piece.render(),
                             f"{name} b{budget}: render drifted from the plan")
        # The step that actually happened to the stack2 run: LightRAG sanitises the whole
        # field, THEN splits on the delimiter. Replay that here for every shape.
        self.assertTrue(is_conservative_marker(delimiter),
                        f"{name} b{budget}: delimiter {delimiter!r} is not made of characters "
                        f"the pipeline cannot touch")
        survives, why = pieces_survive(rendered, delimiter, [p.render() for p in plan.pieces])
        self.assertTrue(survives, f"{name} b{budget}: {why}")
        self.assertNotIn(delimiter, sanitize_like_lightrag(rendered).replace(delimiter, ""),
                         f"{name} b{budget}: a body carries marker-shaped text")
        return plan

    def test_invariants_hold_for_every_fixture_at_three_budgets(self):
        for name, text in FIXTURES.items():
            for budget in (60, 120, 350):
                with self.subTest(fixture=name, budget=budget):
                    self.check(name, text, budget)

    def test_invariants_hold_without_headers(self):
        for name, text in FIXTURES.items():
            with self.subTest(fixture=name):
                plan = self.check(name, text, 120, headers=False)
                self.assertTrue(all(piece.header is None for piece in plan.pieces))

    def test_the_plan_is_deterministic(self):
        for name, text in FIXTURES.items():
            with self.subTest(fixture=name):
                first = chunk_document(text, filename="Doc_2026-01-02.txt",
                                       budget_tokens=140, header_tokens=25)
                second = chunk_document(text, filename="Doc_2026-01-02.txt",
                                        budget_tokens=140, header_tokens=25)
                self.assertEqual([piece.render() for piece in first.pieces],
                                 [piece.render() for piece in second.pieces])

    def test_no_mid_word_cut(self):
        """Splitting a word would change the word multiset; that is what this checks."""
        text = "Supercalifragilistic expialidocious " * 20 + "\n\n" + "hyphen-ated tokens here " * 20
        plan = chunk_document(text, filename="words.txt", budget_tokens=40,
                              header_tokens=10)
        for piece in plan.pieces:
            self.assertNotIn("Supercalifragilisticexpialidocious", piece.body)
        self.assertEqual(Counter(_content_words(" ".join(p.body for p in plan.pieces))),
                         Counter(_content_words(text)))


# ── headers ────────────────────────────────────────────────────────────────

class TestHeaders(unittest.TestCase):
    def test_header_format_and_omitted_fields(self):
        plan = chunk_document(HEADINGS, filename="Procurement_2026-03-15.md",
                              budget_tokens=300, header_tokens=40)
        header = plan.pieces[0].header
        self.assertTrue(header.startswith("[Doc: Procurement | Date: 2026-03-15"),
                        header)
        self.assertTrue(header.endswith("]"), header)
        self.assertIn("Part 1/", header)
        titleless = chunk_document(HEADINGS, filename="transcript.txt",
                                   budget_tokens=300, header_tokens=40)
        self.assertNotIn("Doc:", titleless.pieces[0].header)
        self.assertIn("Part 1/", titleless.pieces[0].header)

    def test_header_is_capped_and_fields_are_dropped_in_priority_order(self):
        deep = ("# Alpha\n## Beta\n## Gamma\n### Delta deep inside the document with more\n"
                "Some body text that belongs to the deepest section of this file.\n")
        plan = chunk_document(deep, filename="Alpha_2026-03-15.md",
                              budget_tokens=300, header_tokens=10)
        for piece in plan.pieces:
            self.assertLessEqual(token_count(piece.header), 10,
                                 f"header over cap: {piece.header!r}")
        self.assertIn("header-fields-dropped-to-fit", plan.warnings)
        last = plan.pieces[-1].header
        self.assertNotIn("Section:", last, "the section path goes first")
        self.assertNotIn("Date:", last, "then the date")
        self.assertIn("Doc: Alpha", last, "the document title is the LAST field dropped")

    def test_an_impossibly_small_cap_still_leaves_a_part_marker(self):
        plan = chunk_document(HEADINGS, filename="Procurement_2026-03-15.md",
                              budget_tokens=300, header_tokens=1)
        for piece in plan.pieces:
            self.assertTrue(piece.header.startswith("[Part ") and piece.header.endswith("]"),
                            piece.header)
        self.assertIn("header-minimised", plan.warnings)

    def test_header_never_eats_the_body_budget(self):
        text = "Word. " * 200
        plan = chunk_document(text, filename="A_Very_Long_Document_Name_For_A_Test_2026.txt",
                              budget_tokens=80, header_tokens=60)
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, 80)

    def test_headers_false_means_no_header_at_all(self):
        plan = chunk_document(HEADINGS, filename="Procurement.md", budget_tokens=120,
                              headers=False)
        self.assertTrue(all(piece.header is None for piece in plan.pieces))
        self.assertEqual(plan.pieces[0].token_estimate,
                         token_count(plan.pieces[0].body))

    def test_shuffled_placebo_keeps_the_words_and_moves_the_attachment(self):
        honest = chunk_document(HEADINGS, filename="Procurement_2026-03-15.md",
                                max_piece_tokens=80, min_piece_tokens=20, header_tokens=40)
        placebo = chunk_document(HEADINGS, filename="Procurement_2026-03-15.md",
                                 max_piece_tokens=80, min_piece_tokens=20, header_tokens=40,
                                 headers="shuffled")
        self.assertEqual([_content_words(piece.body) for piece in honest.pieces],
                         [_content_words(piece.body) for piece in placebo.pieces],
                         "the placebo may not change any body")
        self.assertGreater(len(honest.pieces), 2)
        self.assertTrue(any(piece.section_path for piece in honest.pieces),
                        "this fixture must produce distinct section paths")
        self.assertEqual(Counter(_content_words(" ".join(p.header or "" for p in honest.pieces))),
                         Counter(_content_words(" ".join(p.header or "" for p in placebo.pieces))),
                         "same header words, wrong piece")
        self.assertNotEqual([piece.header for piece in honest.pieces],
                            [piece.header for piece in placebo.pieces],
                            "the attachment has to actually move")
        self.assertEqual(placebo.stats.get("headers_shuffled"), 1)
        for piece in placebo.pieces:
            self.assertIn(f"Part {piece.part_index}/{piece.part_total}", piece.header,
                          "part numbers stay about the piece they now sit on")

    def test_body_and_header_together_are_the_number_the_budget_applies_to(self):
        plan = chunk_document(FURNITURE, filename="Deck.txt", budget_tokens=70,
                              header_tokens=25)
        for piece in plan.pieces:
            self.assertEqual(piece.token_estimate, token_count(piece.render()))


# ── delimiter and the LightRAG params ─────────────────────────────────────

class TestDelimiterAndRender(unittest.TestCase):
    def test_the_chosen_delimiter_is_absent_from_every_body(self):
        plan = chunk_document(TRANSCRIPT, filename="Standup.txt", budget_tokens=80)
        delimiter = pick_delimiter(plan)
        self.assertIn(delimiter, DELIMITER_CANDIDATES)
        self.assertNotIn(delimiter, "".join(piece.body for piece in plan.pieces))

    def test_an_explicit_delimiter_that_occurs_in_the_text_is_refused(self):
        plan = chunk_document(TRANSCRIPT, filename="Standup.txt", budget_tokens=80)
        with self.assertRaises(StructureChunkerError) as caught:
            pick_delimiter(plan, "(Alice Chen)")
        self.assertIn("would cut where this plan did not", str(caught.exception))

    def test_a_whitespace_only_delimiter_is_refused(self):
        plan = chunk_document(PLAIN, filename="notes.txt", budget_tokens=80)
        with self.assertRaises(StructureChunkerError):
            pick_delimiter(plan, "   ")

    def test_params_are_the_shape_lightrag_1_5_7_accepts(self):
        plan = chunk_document(MARKDOWN_TABLE, filename="Costs.md", budget_tokens=90)
        text, params = render_for_lightrag(plan)
        self.assertEqual(params["strategy"], "fixed_token")
        block = params["params"]
        self.assertIs(block["split_by_character_only"], True,
                      "a bool, not 1: that image answers a bool-as-int with a 422")
        self.assertIsInstance(block["chunk_token_size"], int)
        self.assertEqual(block["chunk_token_size"], 90)
        self.assertIsInstance(block["split_by_character"], str)
        # A join, never a wrap. The stored field is stripped before it is split, so a marker
        # at either edge loses its newlines and the payload stops carrying the delimiter the
        # request named - the failure mode that made \x1e useless.
        self.assertFalse(text.startswith(block["split_by_character"]))
        self.assertFalse(text.endswith(block["split_by_character"]))
        self.assertEqual(text.count(block["split_by_character"]), len(plan.pieces) - 1)

    def test_an_empty_plan_cannot_be_rendered(self):
        plan = chunk_document("", filename="blank.txt")
        with self.assertRaises(StructureChunkerError):
            render_for_lightrag(plan)


# ── token counting ─────────────────────────────────────────────────────────

class TestTokenCounting(unittest.TestCase):
    def test_the_estimator_over_counts_never_under_counts_short_text(self):
        for text in ("a", "hello world", "1", "a b c d e f g h i j k l m n o p",
                     "internationalization", "word. " * 30):
            self.assertGreaterEqual(token_count(text), len(text.split()),
                                    f"under-counted: {text!r}")

    def test_which_counter_was_used_is_reported_not_assumed(self):
        plan = chunk_document(PLAIN, filename="notes.txt", budget_tokens=80)
        self.assertIn(plan.stats["tokenizer"],
                      ("tiktoken:cl100k_base", "conservative-overcount"))
        self.assertEqual(plan.stats["tokenizer"], tokenizer_name())

    def test_token_percentiles_are_recorded(self):
        plan = chunk_document(PLAIN, filename="notes.txt", budget_tokens=50)
        self.assertLessEqual(plan.stats["tokens_p50"], plan.stats["tokens_p95"])
        self.assertLessEqual(plan.stats["tokens_p95"], plan.stats["tokens_max"])
        self.assertLessEqual(plan.stats["tokens_max"], 50)



# ── worker wiring (behind STRUCTURE_CHUNKING_ENABLED) ──────────────────────

import asyncio                                        # noqa: E402
import hashlib                                        # noqa: E402
import tempfile                                       # noqa: E402
from unittest.mock import AsyncMock, patch            # noqa: E402

os.environ.setdefault("RAG_API_KEY", "1234")
os.environ.setdefault("LIGHTRAG_INTERNAL_API_KEY", "internal")
os.environ.setdefault("LIGHTRAG_BASE_URL", "http://example:9621")
os.environ.setdefault("SOURCE_DOCS_DIR", "/tmp/source_docs")
os.environ.setdefault("UPLOADS_DIR", "/tmp/uploads")
os.environ.setdefault("STATE_DB_PATH", "/tmp/state.db")

from typing import Any                                # noqa: E402

import app.worker as worker                           # noqa: E402
from app.pdf_extract import pdftotext_available      # noqa: E402
from tests.test_document_extract import make_docx, make_pdf  # noqa: E402
import chunk_embed_bench as bench                     # noqa: E402


def _run_async(coroutine):
    return asyncio.new_event_loop().run_until_complete(coroutine)


class _FakeStore:
    """The five state-store calls submit_file makes, in memory."""

    def __init__(self) -> None:
        self.docs: dict[str, dict] = {}

    async def get_document_by_path(self, path: str) -> dict | None:
        return self.docs.get(path)

    async def get_documents_by_sha256(self, digest: str) -> list[dict]:
        return [doc for doc in self.docs.values() if doc.get("sha256") == digest]

    async def upsert_document(self, path: str, **fields) -> None:
        self.docs.setdefault(path, {"path": path}).update(fields)

    async def update_document_state(self, path: str, **fields) -> bool:
        self.docs.setdefault(path, {"path": path}).update(fields)
        return True

    async def increment_retry(self, path: str, **kwargs) -> int:
        return 0

    def meta(self, path: str) -> dict:
        block = json.loads(self.docs.get(path, {}).get("warnings_json") or "{}").get("worker")
        return block if isinstance(block, dict) else {}


DOC_TEXT = ("Quarterly numbers\n\n| Region | Q1 | Q2 |\n|---|---|---|\n"
            + "\n".join(f"| north {index} | {index * 3} | {index * 5} |"
                        for index in range(1, 8))
            + "\n\nThe north grew every quarter of the year, on the same headcount.\n")


class TestWorkerWiring(unittest.TestCase):
    """Off must be byte-identical; on must send the params; failure must fall back."""

    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp())
        (self.root / "source_docs").mkdir(parents=True)
        self.store = _FakeStore()
        self.post_json = AsyncMock(return_value={"track_id": "trk-1"})
        self.post_file = AsyncMock(return_value={"track_id": "trk-1"})
        patcher = patch.multiple(
            worker.settings,
            source_docs_dir=str(self.root / "source_docs"),
            uploads_dir=str(self.root / "uploads"),
            transcript_prep_enabled=False,
            structure_chunking_enabled=False,
            structure_chunking_max_piece_tokens=70,
            structure_chunking_min_piece_tokens=20,
            structure_chunking_header_tokens=20,
            structure_chunking_headers=True,
        )
        patcher.start()
        self.addCleanup(patcher.stop)

    def _patches(self):
        return patch.multiple(
            worker,
            get_document_by_path=self.store.get_document_by_path,
            get_documents_by_sha256=self.store.get_documents_by_sha256,
            update_document_state=self.store.update_document_state,
            upsert_document=self.store.upsert_document,
            increment_retry=self.store.increment_retry,
        )

    def _submit(self, name: str, content: str) -> dict:
        path = self.root / "source_docs" / name
        path.write_text(content, encoding="utf-8")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        self.store.docs[str(path)] = {"path": str(path), "sha256": digest,
                                      "status": "accepted", "company": None,
                                      "filename": name, "original_filename": name,
                                      "warnings_json": None}
        with self._patches(), patch.object(worker.client, "post_json", self.post_json), \
             patch.object(worker.client, "post_file", self.post_file):
            _run_async(worker.submit_file(path, "filesystem", digest))
        body = self.post_json.await_args.args[1]
        return {"path": str(path), "digest": digest, "body": body, "text": body["text"]}

    def test_flag_off_body_is_byte_identical_to_today(self):
        got = self._submit("numbers.md", DOC_TEXT)
        self.assertEqual(set(got["body"]), {"text", "file_source"},
                         "no chunking key at all when the flag is off")
        self.assertEqual(got["text"], worker.normalize_text_submission_payload(DOC_TEXT))
        self.assertNotIn("structure_chunking", self.store.meta(got["path"]))

    def test_flag_on_posts_the_delimiter_params_and_keeps_the_dedupe_hash(self):
        with patch.object(worker.settings, "structure_chunking_enabled", True):
            got = self._submit("numbers_2026-04.md", DOC_TEXT)
        params = got["body"]["chunking"]
        self.assertEqual(params["strategy"], "fixed_token")
        self.assertTrue(params["params"]["split_by_character_only"])
        # chunk_token_size is NOT the budget: it must clear the largest piece, or LightRAG
        # is free to re-cut the pieces we chose (see the fit-check tests below).
        self.assertGreaterEqual(params["params"]["chunk_token_size"], 120)
        self.assertIn(params["params"]["split_by_character"], got["text"])
        self.assertIn("[Doc: numbers 2026 04 | Part 1/", got["text"],
                      "the header is built from the file name, title only here")
        self.assertIn("Part 2/", got["text"],
                      "the cap is below the document's natural units, so it cuts")
        row = self.store.docs[got["path"]]
        self.assertEqual(row["sha256"], got["digest"],
                         "the dedupe hash stays the ORIGINAL file bytes")

    def test_flag_on_records_provenance_as_counts_and_names_only(self):
        with patch.object(worker.settings, "structure_chunking_enabled", True):
            got = self._submit("numbers_2026-04.md", DOC_TEXT)
        meta = self.store.meta(got["path"])
        self.assertEqual(meta["structure_chunking"], "applied")
        self.assertEqual(meta["structure_chunking_kind"], "table")
        self.assertGreater(meta["structure_chunking_pieces"], 1)
        self.assertLessEqual(meta["structure_chunking_tokens_max"], 120)
        self.assertIn(meta["structure_chunking_tokenizer"],
                      ("tiktoken:cl100k_base", "conservative-overcount"))
        self.assertIsInstance(meta["structure_chunking_warnings"], list)
        blob = json.dumps(meta)
        for leak in ("Region", "north 1", "headcount"):
            self.assertNotIn(leak, blob, "provenance must carry no document text")

    def test_headers_false_reaches_the_plan(self):
        with patch.multiple(worker.settings, structure_chunking_enabled=True,
                            structure_chunking_headers=False):
            got = self._submit("numbers_2026-04.md", DOC_TEXT)
        self.assertNotIn("[Doc:", got["text"])
        self.assertIn("split_by_character", got["body"]["chunking"]["params"])

    def test_a_failing_planner_falls_back_to_the_untouched_text(self):
        def explode(*args, **kwargs):
            raise RuntimeError("planner exploded")
        with patch.object(worker.settings, "structure_chunking_enabled", True), \
             patch.object(worker, "chunk_document", explode):
            got = self._submit("numbers_2026-04.md", DOC_TEXT)
        self.assertEqual(got["text"], worker.normalize_text_submission_payload(DOC_TEXT),
                         "the document still goes in, unchunked")
        self.assertNotIn("chunking", got["body"], "and no params are sent with old text")
        meta = self.store.meta(got["path"])
        self.assertEqual(meta["structure_chunking"], "failed")
        self.assertIn("planner exploded", meta["structure_chunking_reason"])

    def test_binary_upload_path_is_untouched(self):
        path = self.root / "source_docs" / "deck.pdf"
        path.write_bytes(b"%PDF-1.4 fake")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        self.store.docs[str(path)] = {"path": str(path), "sha256": digest,
                                      "status": "accepted", "company": None,
                                      "filename": "deck.pdf", "original_filename": "deck.pdf",
                                      "warnings_json": None}
        with patch.object(worker.settings, "structure_chunking_enabled", True), \
             self._patches(), patch.object(worker.client, "post_json", self.post_json), \
             patch.object(worker.client, "post_file", self.post_file):
            _run_async(worker.submit_file(path, "filesystem", digest))
        self.assertTrue(self.post_file.await_count == 1,
                        "a .pdf must still go through /documents/upload")
        self.assertEqual(self.post_json.await_count, 0,
                         "the extractor route is a separate decision, not this one")


class TestBinaryRoutes(unittest.TestCase):
    """PDF, DOCX and PPTX reach LightRAG pre-cut, or they reach it as they do today.

    Every failure mode has to land on the upload route with a recorded reason: a parser bug
    may change how a document is chunked, never whether it is ingested.
    """

    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp())
        self.docs = self.root / "source_docs"
        self.docs.mkdir(parents=True)
        self.store = _FakeStore()
        self.post_json = AsyncMock(return_value={"track_id": "trk-1"})
        self.post_file = AsyncMock(return_value={"track_id": "trk-1"})
        patcher = patch.multiple(
            worker.settings,
            source_docs_dir=str(self.docs), uploads_dir=str(self.root / "uploads"),
            transcript_prep_enabled=False, structure_chunking_enabled=True,
            structure_chunking_max_piece_tokens=160,
            structure_chunking_min_piece_tokens=20, structure_chunking_header_tokens=20,
            structure_chunking_headers=True,
            structure_chunking_fit_margin_tokens=64,
            structure_chunking_embed_window_tokens=8192,
            structure_chunking_embed_tokenizer_path="",
            structure_chunking_fail_on_lost_text=False,
            worker_pptx_extract_enabled=True, worker_pdf_extract_enabled=True,
            worker_docx_extract_enabled=True, worker_pdf_text_backend="auto",
            worker_pdf_max_bytes=0, worker_pdf_max_chars=0, worker_docx_max_bytes=0,
            worker_docx_max_chars=0)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _submit(self, path: Path) -> dict[str, Any]:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        self.store.docs[str(path)] = {"path": str(path), "sha256": digest, "status": "accepted",
                                      "company": None, "filename": path.name,
                                      "original_filename": path.name, "warnings_json": None}
        with patch.multiple(worker,
                            get_document_by_path=self.store.get_document_by_path,
                            get_documents_by_sha256=self.store.get_documents_by_sha256,
                            update_document_state=self.store.update_document_state,
                            upsert_document=self.store.upsert_document,
                            increment_retry=self.store.increment_retry), \
             patch.object(worker.client, "post_json", self.post_json), \
             patch.object(worker.client, "post_file", self.post_file):
            _run_async(worker.submit_file(path, "filesystem", digest))
        return {"path": str(path), "digest": digest,
                "json_body": self.post_json.await_args.args[1] if self.post_json.await_count else None,
                "text_route": self.post_json.await_count, "upload_route": self.post_file.await_count,
                "meta": self.store.meta(str(path))}

    def test_a_docx_table_goes_in_pre_cut_with_the_header_repeated(self):
        rows = [["Region", "Q1", "Q2"]] + [
            [f"north {index}", str(index * 3), str(index * 5)] for index in range(1, 61)]
        path = make_docx([("Heading1", "Budget detail"), ("", "Every region, every quarter.")],
                         [rows], ["p", "p", "t"], self.docs / "budget_2026-04.docx")
        got = self._submit(path)
        self.assertEqual((got["upload_route"], got["text_route"]), (0, 1),
                         "the worker parsed it, so the text route is the one that carries cuts")
        body = got["json_body"]
        params = body["chunking"]
        self.assertTrue(params["params"]["split_by_character_only"])
        delimiter = params["params"]["split_by_character"]
        pieces = [piece for piece in body["text"].split(delimiter) if piece.strip()]
        self.assertGreater(len(pieces), 1)
        table_pieces = [piece for piece in pieces if "| north" in piece]
        self.assertGreaterEqual(len(table_pieces), 2, "the 60-row table had to split")
        for piece in table_pieces:
            self.assertIn("Region | Q1 | Q2", piece,
                          "every piece the table split across keeps the header row")
        meta = got["meta"]
        self.assertEqual(meta["structure_extract"], "worker")
        self.assertEqual(meta["structure_extract_kind"], "docx")
        self.assertEqual(meta["structure_chunking"], "applied")
        self.assertEqual(meta["structure_chunking_kind"], "table")
        self.assertLessEqual(meta["structure_chunking_tokens_max"], 160)
        self.assertEqual(meta["structure_chunking_lost_word_tokens"], 0)
        self.assertEqual(got["digest"], hashlib.sha256(path.read_bytes()).hexdigest(),
                         "the dedupe hash is still the file bytes, not the parsed text")

    def test_a_pdf_loses_its_page_furniture_but_keeps_every_row(self):
        if not pdftotext_available():
            self.skipTest("poppler's pdftotext is not installed here")
        pages = [
            [(72, 760, "Acme Corporation confidential"),
             (72, 720, "1. Revenue"),
             (72, 700, "Region        Q1         Q2"),
             (72, 686, "north         120        150"),
             (180, 672, "grew on the same headcount"),
             (72, 658, "south         70         90")],
            [(72, 760, "Acme Corporation confidential"),
             (72, 720, "2. Costs"),
             (72, 700, "Costs fell in the north region after the vendor renegotiation.")],
            [(72, 760, "Acme Corporation confidential"),
             (72, 720, "3. Risks"),
             (72, 700, "Risks were unchanged from the previous quarter, per the register.")],
        ]
        path = make_pdf(pages, self.docs / "report_2026-04.pdf")
        got = self._submit(path)
        self.assertEqual((got["upload_route"], got["text_route"]), (0, 1))
        text = got["json_body"]["text"]
        self.assertNotIn("Acme Corporation confidential", text,
                         "the running header is furniture and is counted, not submitted")
        self.assertIn("grew on the same headcount", text, "the wrapped row stayed one row")
        meta = got["meta"]
        self.assertEqual(meta["structure_extract_kind"], "pdf")
        self.assertGreaterEqual(meta["structure_chunking_dropped_lines"], 3)

    def test_a_deck_takes_the_same_preparation_as_a_text_file(self):
        try:
            from pptx import Presentation
        except ImportError:
            self.skipTest("python-pptx is not installed")
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Regions"
        table = deck.slides[0].shapes.add_table(3, 3, 100000, 2000000, 4000000, 1000000).table
        for column, heading in enumerate(("Region", "Q1", "Q2")):
            table.cell(0, column).text = heading
        for row, values in enumerate((("north", "12", "15"), ("south", "7", "9")), start=1):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        path = self.docs / "deck.pptx"
        deck.save(str(path))
        got = self._submit(path)
        self.assertEqual((got["upload_route"], got["text_route"]), (0, 1))
        self.assertIn("chunking", got["json_body"],
                      "the deck is pre-cut like every other document now")

    def test_the_pdf_backend_that_ran_is_on_the_row(self):
        import sys
        import types
        import unittest.mock as mock
        page = types.SimpleNamespace(extract_text=lambda: "1. Revenue\nnorth   120   150\n"
                                     "the north region grew on the same headcount while the "
                                     "carrier kept the lane open for the whole term again.")
        reader = types.SimpleNamespace(pages=[page], is_encrypted=False, decrypt=lambda _: True)
        pypdf = types.ModuleType("pypdf")
        pypdf.PdfReader = lambda path: reader
        with patch.object(worker.settings, "worker_pdf_text_backend", "pypdf"), \
             mock.patch.dict(sys.modules, {"pypdf": pypdf}), \
             mock.patch("app.pdf_extract.pdftotext_available", lambda: False):
            make_pdf([[(72, 760, "1. Revenue")]], self.docs / "flat_2026-04.pdf")
            got = self._submit(self.docs / "flat_2026-04.pdf")
        meta = got["meta"]
        self.assertEqual(meta["structure_extract"], "worker")
        self.assertEqual(meta["structure_extract_backend"], "pypdf",
                         "the row has to say which reader produced the text")

    def test_with_the_flag_off_a_pdf_is_still_uploaded_untouched(self):
        path = self.docs / "plain.pdf"
        path.write_bytes(b"%PDF-1.4\nnot really, but the route never reads it\n")
        with patch.object(worker.settings, "structure_chunking_enabled", False):
            got = self._submit(path)
        self.assertEqual((got["upload_route"], got["text_route"]), (1, 0))
        self.assertNotIn("structure_extract", got["meta"], "no new metadata either")

    def test_a_kill_switch_on_one_type_leaves_the_other_alone(self):
        rows = [["Region", "Q1"], ["north", "12"]]
        docx = make_docx([("", "table")], [rows], ["p", "t"], self.docs / "one.docx")
        with patch.object(worker.settings, "worker_docx_extract_enabled", False):
            got = self._submit(docx)
        self.assertEqual((got["upload_route"], got["text_route"]), (1, 0))

    def test_a_scan_falls_back_to_upload_with_the_reason_recorded(self):
        path = make_pdf([[], []], self.docs / "scan.pdf")
        got = self._submit(path)
        self.assertEqual((got["upload_route"], got["text_route"]), (1, 0),
                         "the document still gets ingested; only the cuts are lost")
        meta = got["meta"]
        self.assertEqual(meta["structure_extract"], "fallback_upload")
        self.assertEqual(meta["structure_extract_error"], "pdf_empty")
        self.assertIn("no extractable text", meta["structure_extract_reason"])

    def test_a_corrupt_docx_falls_back_to_upload(self):
        path = self.docs / "broken.docx"
        path.write_bytes(b"definitely not a zip")
        got = self._submit(path)
        self.assertEqual((got["upload_route"], got["text_route"]), (1, 0))
        self.assertEqual(got["meta"]["structure_extract_error"], "docx_extract_failed")

    def test_a_piece_over_the_ceiling_is_not_submitted_chunked(self):
        path = self.docs / "monster.txt"
        path.write_text("start " + "x" * 1200 + " end\n", encoding="utf-8")
        with patch.object(worker.settings, "structure_chunking_max_piece_tokens", 200):
            got = self._submit(path)
        self.assertEqual(got["meta"]["structure_chunking"], "fallback")
        self.assertIn("MAX_PIECE_TOKENS", got["meta"]["structure_chunking_reason"])
        self.assertNotIn("chunking", got["json_body"],
                         "an uncut submission beats a submission that breaks the ceiling")

    def test_the_window_check_uses_the_embedders_own_token_count_when_configured(self):
        tokenizer_dir = self.root / "tokenizer"
        tokenizer_dir.mkdir()
        (tokenizer_dir / "vocab.txt").write_text("[CLS]\n[SEP]\n[UNK]\nthe\n", encoding="utf-8")
        path = self.docs / "numbers.txt"
        path.write_text("the " * 40 + "\n", encoding="utf-8")
        with patch.multiple(worker.settings,
                            structure_chunking_embed_tokenizer_path=str(tokenizer_dir),
                            structure_chunking_embed_window_tokens=30):
            got = self._submit(path)
        meta = got["meta"]
        self.assertEqual(meta["structure_chunking"], "fallback")
        self.assertIn("embedder window", meta["structure_chunking_reason"])
        self.assertIn("embedder tokens", meta["structure_chunking_fit_basis"])
        self.assertGreater(meta["structure_chunking_embedder_tokens_max"], 30)
        self.assertNotIn("chunking", got["json_body"])

    def test_the_default_window_check_runs_on_the_conservative_estimate(self):
        path = self.docs / "long.txt"
        path.write_text("word " * 200 + "\n", encoding="utf-8")
        with patch.object(worker.settings, "structure_chunking_embed_window_tokens", 120):
            got = self._submit(path)
        self.assertEqual(got["meta"]["structure_chunking"], "fallback")
        self.assertIn("estimate", got["meta"]["structure_chunking_fit_basis"])

    def test_chunk_token_size_clears_the_largest_piece_so_nothing_is_re_cut(self):
        path = self.docs / "table.txt"
        lines = ["| Region | Q1 | Q2 |", "|---|---|---|"] + [
            f"| north {index} | {index * 3} | {index * 5} |" for index in range(1, 40)]
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        got = self._submit(path)
        params = got["json_body"]["chunking"]["params"]
        largest = got["meta"]["structure_chunking_tokens_max"]
        self.assertGreaterEqual(params["chunk_token_size"], largest + 64,
                                "if chunk_token_size sat at the budget, LightRAG could re-cut")

    def test_a_payload_whose_cuts_do_not_round_trip_is_not_submitted_chunked(self):
        real = worker.render_for_lightrag

        def markers_removed(plan, delimiter=None, *, chunk_token_size=None):
            text, params = real(plan, delimiter, chunk_token_size=chunk_token_size)
            marker = params["params"]["split_by_character"]
            return text.replace(marker, "\n"), params

        path = self.docs / "wide.txt"
        path.write_text("\n".join(f"## Section {index}\n\n" + " ".join(
            f"carrier{index} lane depot tariff customs origin papers margin"
            for _ in range(120)) for index in range(12)) + "\n", encoding="utf-8")
        with patch.object(worker, "render_for_lightrag", markers_removed):
            got = self._submit(path)
        self.assertEqual(got["meta"]["structure_chunking"], "fallback")
        self.assertIn("cut markers", got["meta"]["structure_chunking_reason"])
        self.assertLess(got["meta"]["structure_chunking_delimiters"],
                        got["meta"]["structure_chunking_pieces"])
        self.assertNotIn("chunking", got["json_body"],
                         "no params without the matching markers in the text")

    def test_a_delimiter_the_sanitiser_strips_is_refused_before_submit(self):
        # The stack2 failure, reproduced offline: \x1e passes through the API and is deleted
        # by LightRAG's own input sanitiser, so the document arrives as one segment. The
        # pre-submit replay is what catches it now.
        def old_way(plan, delimiter=None, *, chunk_token_size=None):
            # join-only, so the marker count is right and only the survival replay can stop it
            body = "\x1e".join(piece.render() for piece in plan.pieces)
            return body, {"strategy": "fixed_token",
                          "params": {"chunk_token_size": 400, "split_by_character": "\x1e",
                                     "split_by_character_only": True}}

        path = self.docs / "wide.txt"
        path.write_text("\n".join(f"## Section {index}\n\n" + " ".join(
            f"carrier{index} lane depot tariff customs origin papers margin"
            for _ in range(120)) for index in range(12)) + "\n", encoding="utf-8")
        with patch.object(worker, "render_for_lightrag", old_way):
            got = self._submit(path)
        self.assertEqual(got["meta"]["structure_chunking"], "fallback")
        self.assertIn("sanitiser", got["meta"]["structure_chunking_reason"])
        self.assertNotEqual(got["meta"]["structure_chunking_sanitiser_check"], "ok",
                            "the reason is recorded, not just the fact of the refusal")
        self.assertNotIn("chunking", got["json_body"])

    def test_control_characters_are_taken_out_at_plan_time_and_reported(self):
        # A form feed rides in on a page break and a NUL rides in on a bad decode. Both are
        # invisible in the stored text and both move where a piece reads as ending, so the
        # chunker normalises them BEFORE planning (planned text == stored text) and names them.
        filler = " ".join(f"carrier{index} lane depot tariff customs" for index in range(120))
        path = self.docs / "paged.txt"
        path.write_text("page one\x0ctop line\n\n" + filler + "\n\n\x0cpage two\x00tail\n",
                        encoding="utf-8")
        got = self._submit(path)
        self.assertEqual(got["meta"]["structure_chunking"], "applied")
        report = got["meta"]["structure_chunking_stripped_chars"]
        self.assertIn("U+000C", report)
        self.assertIn("U+0000", report)
        posted = got["json_body"]["text"]
        for ch in ("\x0c", "\x00"):
            self.assertNotIn(ch, posted, f"U+{ord(ch):04X} reached the request")
        self.assertEqual(got["meta"].get("structure_chunking_sanitiser_check"), "ok")

    def test_the_submitted_chunk_token_size_clears_every_segment_it_will_be_handed(self):
        # doc08 in the stack2 run: chunk_token_size 2579 against a segment LightRAG counted
        # at 2592. Under split_by_character_only that is a rejection, not a re-cut, so the
        # number goes out in the counter LightRAG counts with and clears the payload.
        filler = " ".join(f"carrier{index} lane depot tariff customs" for index in range(120))
        path = self.docs / "wide.txt"
        path.write_text(filler + "\n\n" + filler.replace("carrier", "depot") + "\n",
                        encoding="utf-8")
        def cl100k_shaped(text):                    # ~2 tokens per word: higher than the
            return max(1, 2 * len(text.split()))    # over-count's word rung on this prose
        with patch.object(sc, "_TOKENIZER", ("tiktoken:fake", cl100k_shaped)), \
             patch.object(worker.settings, "structure_chunking_max_piece_tokens", 900):
            got = self._submit(path)
        self.assertEqual(got["meta"]["structure_chunking"], "applied")
        self.assertEqual(got["meta"]["structure_chunking_chunk_token_basis"], "tiktoken:fake")
        params = got["json_body"]["chunking"]["params"]
        marker = params["split_by_character"]
        segments = [seg for seg in got["json_body"]["text"].split(marker) if seg.strip()]
        self.assertGreater(len(segments), 1)
        margin = int(worker.settings.structure_chunking_fit_margin_tokens or 0)
        biggest = max(cl100k_shaped(seg) for seg in segments)
        self.assertGreaterEqual(params["chunk_token_size"], biggest + margin,
                                f"cts {params['chunk_token_size']} would reject a segment "
                                f"the counter names at {biggest}")

    def test_a_coverage_gap_is_recorded_and_only_blocks_when_told_to(self):
        path = self.docs / "words.txt"
        path.write_text("alpha beta gamma delta epsilon zeta eta theta\n", encoding="utf-8")
        def leaking(source_text, plan):
            return {"source_word_tokens": 9, "piece_word_tokens": 9,
                    "furniture_word_tokens": 0, "missing_word_tokens": 3,
                    "unaccounted_word_tokens": 3, "unaccounted_distinct": 3,
                    "covered_pct": 66.67}
        with patch.object(worker, "coverage_gaps", leaking):
            got = self._submit(path)
            self.assertEqual(got["meta"]["structure_chunking_lost_word_tokens"], 3)
            self.assertEqual(got["meta"]["structure_chunking"], "applied",
                             "a gap is recorded by default: blocking ingest on a heuristic "
                             "is a worse outcome than the gap")
            with patch.object(worker.settings, "structure_chunking_fail_on_lost_text", True):
                strict = self._submit(path)
        self.assertEqual(strict["meta"]["structure_chunking"], "fallback")
        self.assertIn("3 word tokens", strict["meta"]["structure_chunking_reason"])
        self.assertNotIn("chunking", strict["json_body"])

    def test_transcript_text_is_still_cut_at_turn_boundaries(self):
        path = self.docs / "meeting.txt"
        turns = [f"Alice: point number {index} about the quarterly budget review meeting."
                 for index in range(1, 15)]
        path.write_text("\n".join(turns) + "\n", encoding="utf-8")
        got = self._submit(path)
        meta = got["meta"]
        self.assertEqual(meta["structure_chunking_kind"], "transcript")
        delimiter = got["json_body"]["chunking"]["params"]["split_by_character"]
        pieces = [piece for piece in got["json_body"]["text"].split(delimiter) if piece.strip()]
        self.assertGreater(len(pieces), 1)
        for piece in pieces:
            self.assertTrue(piece.splitlines()[1].startswith("Alice:"),
                            "a turn is never split, so a piece starts at a turn")


# ── the benchmark arm, so the same function is measured and shipped ────────

class TestBenchmarkStrategy(unittest.TestCase):
    def test_the_grid_carries_the_structure_arms(self):
        configs = bench.load_configs(CONFIGS)
        by_name = {config["name"]: config for config in configs}
        for name in ("structure-700", "structure-700-nohdr", "structure-700-shuffled",
                     "structure-1500", "structure-1500-nohdr", "structure-1500-shuffled"):
            self.assertIn(name, by_name, f"the grid is missing {name}")
        self.assertEqual(by_name["structure-1500"]["chunker"]["strategy"], "structure")
        self.assertEqual(by_name["structure-1500"]["chunker"]["max_piece_tokens"], 1500)
        self.assertEqual(by_name["structure-1500"]["chunker"]["min_piece_tokens"], 150)
        self.assertIs(by_name["structure-1500-nohdr"]["chunker"]["headers"], False)
        self.assertEqual(by_name["structure-1500-shuffled"]["chunker"]["headers"], "shuffled")
        self.assertNotIn("mxbai-embed-large-v1",
                         {config["embedder"]["model"] for config in configs},
                         "the grid must not re-introduce the retired embedder")

    def test_the_structure_strategy_routes_to_the_shipped_module(self):
        tokenizer = bench.WhitespaceTokenizer()
        pieces = bench.chunk_document(HEADINGS, {"strategy": "structure",
                                                "budget_tokens": 60, "headers": True},
                                      tokenizer)
        self.assertGreater(len(pieces), 1)
        self.assertIn("[Section:", pieces[0])
        from app import structure_chunker as shipped
        direct = shipped.chunk_document(HEADINGS, filename="", budget_tokens=60)
        self.assertEqual(pieces, [piece.render() for piece in direct.pieces],
                         "the benchmark must not keep its own imitation of the chunker")

    def test_the_placebo_is_accepted_by_the_loader_and_moves_headers(self):
        specs = {"headers": "shuffled", "budget_tokens": 60}
        plain = bench.chunk_document(HEADINGS, {"strategy": "structure", "headers": True},
                                     bench.WhitespaceTokenizer())
        shuffled = bench.chunk_document(HEADINGS, {"strategy": "structure", **specs},
                                        bench.WhitespaceTokenizer())
        self.assertNotEqual([chunk.splitlines()[0] for chunk in plain],
                            [chunk.splitlines()[0] for chunk in shuffled])

    def test_bad_structure_settings_are_refused_before_any_embedding(self):
        base = {"name": "a", "chunker": {"strategy": "structure"},
                "embedder": {"model": "m"}}
        for bad in ({"budget_tokens": 5}, {"budget_tokens": "350"},
                    {"header_tokens": -1}, {"headers": "yes"}):
            with self.subTest(setting=bad):
                payload = dict(base, chunker={"strategy": "structure", **bad})
                path = Path(tempfile.mkdtemp()) / "configs.json"
                path.write_text(json.dumps({"configs": [payload]}), encoding="utf-8")
                with self.assertRaises(bench.BenchError):
                    bench.load_configs(path)

    def test_the_error_message_lists_the_structure_strategy(self):
        with self.assertRaises(bench.BenchError) as caught:
            bench.chunk_document("text", {"strategy": "custom"}, bench.WhitespaceTokenizer())
        self.assertIn("structure", str(caught.exception))




class TestWindowNotACap(unittest.TestCase):
    """The shipped defaults: the window is the ceiling, the floor is the only target."""

    def test_the_constants_say_what_the_operator_asked_for(self):
        from app.config import settings
        from app.embed_tokenizer import EMBED_WINDOW_DEFAULT
        from app.structure_chunker import (DEFAULT_MAX_PIECE_TOKENS,
                                          DEFAULT_MIN_PIECE_TOKENS, DEFAULT_MERGE_UNTIL)
        self.assertEqual(EMBED_WINDOW_DEFAULT, 8192)
        self.assertEqual(int(settings.structure_chunking_embed_window_tokens), 8192)
        self.assertLess(DEFAULT_MAX_PIECE_TOKENS, EMBED_WINDOW_DEFAULT,
                        "the piece window is the embedder window minus a safety margin")
        self.assertEqual(DEFAULT_MIN_PIECE_TOKENS, 150, "a floor, not a size target")
        self.assertEqual(DEFAULT_MERGE_UNTIL, "floor")
        self.assertEqual(int(settings.structure_chunking_min_piece_tokens), 150)
        self.assertEqual(int(settings.structure_chunking_max_piece_tokens),
                         DEFAULT_MAX_PIECE_TOKENS)


class TestTranscriptLabelShapes(unittest.TestCase):
    """Who speaks, recognised by how OFTEN a label repeats - not by how the file wraps.

    The corpus contains minutes written as one line per agenda item and transcripts
    labelled "(Name)" or "Name 12:34". A line-anchored parser calls all of those prose,
    and prose is what gets no speaker turn boundaries at all.
    """

    WINDOW = 600
    FLOOR = 120

    def plan(self, text: str, **kwargs):
        return chunk_document(text, filename="Minutes_2026-04.txt",
                              max_piece_tokens=self.WINDOW,
                              min_piece_tokens=self.FLOOR, **kwargs)

    def minutes_line(self, vocabularies: list[list[str]], turns: int = 12) -> str:
        parts: list[str] = []
        for index, topic in enumerate(vocabularies):
            for turn in range(turns):
                word = topic[turn % len(topic)]
                other = topic[(turn + 1) % len(topic)]
                speaker = "Ana" if index % 2 == 0 else "Liang"
                parts.append(f"{speaker}: the {word} position sits against {other} for "
                             f"period {turn}, and the {word} margin moved twice while the "
                             f"{other} desk kept the same note.")
        return " ".join(parts)

    def test_minutes_written_as_one_line_are_a_transcript_not_one_great_paragraph(self):
        text = self.minutes_line([["freight", "carrier", "lane"],
                                  ["tariff", "customs", "origin"]])
        plan = self.plan(text)
        self.assertEqual(plan.kind_detected, "transcript")
        self.assertGreater(len(plan.pieces), 1)
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, self.WINDOW)

    def test_splitting_the_line_loses_no_word(self):
        text = self.minutes_line([["freight", "carrier"], ["tariff", "customs"]])
        plan = self.plan(text)
        gaps = coverage_gaps(text, plan)
        self.assertEqual(gaps["unaccounted_word_tokens"], 0, gaps)

    def test_parenthesised_labels_with_a_colon_are_turns(self):
        lines = []
        for index in range(14):
            name = "Ana" if index % 2 == 0 else "Liang"
            lines.append(f"({name}): the depot carrier lane {index} was reviewed and the "
                         f"freight margin for period {index} moved twice in the same week.")
        plan = self.plan("\n".join(lines))
        self.assertEqual(plan.kind_detected, "transcript")
        self.assertIn("Ana", [piece for piece in plan.pieces] and
                      [name for name in ("Ana", "Liang")
                       if name in "\n".join(piece.body for piece in plan.pieces)])

    def test_a_name_and_timestamp_label_is_a_turn(self):
        lines = []
        for index in range(14):
            name = "Ana" if index % 2 == 0 else "Liang"
            lines.append(f"{name} {index:02d}:{index:02d}")
            lines.append(f"The depot carrier lane {index} was reviewed and the freight "
                         f"margin for period {index} moved twice in the same week, and the "
                         f"customs desk kept the same note about the origin papers.")
        plan = self.plan("\n".join(lines))
        self.assertEqual(plan.kind_detected, "transcript")

    def test_prose_that_uses_colons_freely_is_still_prose(self):
        paragraphs = []
        for index in range(12):
            paragraphs.append(
                f"The report for period {index}: the freight numbers were flat, and the "
                f"review: nothing changed in the depot lane, despite the note: we asked "
                f"twice about the carrier tariff, and the answer: the customs desk holds "
                f"the origin papers until the broker files them again this quarter.")
        plan = self.plan("\n\n".join(paragraphs))
        self.assertNotEqual(plan.kind_detected, "transcript")

    def test_the_inline_splitter_adds_no_fake_turns_to_a_key_value_list(self):
        # KNOWN, REPORTED: a capitalised key/value list ("Vendor: depot 0. Cost: 1200.") is
        # already read as a one-voice transcript by transcript_prep's long-standing bars,
        # and this chunker mirrors those bars on purpose. That pre-existing false positive
        # is the lead's call to change; what this test pins is that the new mid-line label
        # shape does not make it worse - no extra turns, and no invented second voice.
        lines = [f"Vendor: depot {index}. Cost: {1000 + index}. Term: 24 months of service "
                 f"at the same site, and the vendor: the carrier keeps the lane open."
                 for index in range(12)]
        text = "\n".join(lines)
        plan = self.plan(text)
        # No word is invented or duplicated by the label scan, and none is lost.
        self.assertEqual(len(body_words(plan)), len(text.split()))
        self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, self.WINDOW)




class TestHermesTranscriptShapes(unittest.TestCase):
    """The five shapes a real run could not classify. Invented text; shapes only.

    What the general name grammar cannot know, the file says: two labels heading the
    lines, strictly alternating, is a transcript whatever the words are. Stage
    directions alternate just as happily and are still not people.
    """

    WINDOW = 600
    FLOOR = 120

    def plan(self, text: str, **kwargs):
        kwargs.setdefault("max_piece_tokens", self.WINDOW)
        kwargs.setdefault("min_piece_tokens", self.FLOOR)
        return chunk_document(text, filename="Call_2026-04-11.txt", **kwargs)

    def alternating(self, first: str, second: str, lines: int = 40,
                    words: int = 18) -> str:
        out: list[str] = []
        for index in range(lines):
            label = first if index % 2 == 0 else second
            topic = ["freight", "tariff", "carrier", "customs", "depot"][index % 5]
            filler = " ".join(f"{topic}_{index}_{word}" for word in range(words))
            out.append(f"({label}) {filler}")
        return "\n".join(out) + "\n"

    def test_a_pronoun_pair_that_alternates_is_a_transcript(self):
        text = self.alternating("They", "You")
        plan = self.plan(text)
        self.assertEqual(plan.kind_detected, "transcript")
        self.assertGreater(len(plan.pieces), 1)
        gaps = coverage_gaps(text, plan)
        self.assertEqual(gaps["unaccounted_word_tokens"], 0, gaps)

    def test_a_three_letter_and_a_five_letter_label_pair(self):
        for pair in (("They", "You"), ("You", "There"), ("She", "Their")):
            with self.subTest(labels=pair):
                plan = self.plan(self.alternating(*pair, lines=24))
                self.assertEqual(plan.kind_detected, "transcript")

    def test_a_real_name_paired_with_a_pronoun_parses_wholly(self):
        text = self.alternating("Ana", "They", lines=24)
        plan = self.plan(text)
        self.assertEqual(plan.kind_detected, "transcript")
        bodies = "\n".join(piece.body for piece in plan.pieces)
        self.assertIn("(Ana)", bodies)
        self.assertIn("(They)", bodies)

    def test_stage_directions_are_never_rescued_by_alternation(self):
        for pair in (("Cross-talk", "Inaudible"), ("Background noise", "Silence"),
                     ("Laughter", "Applause")):
            with self.subTest(labels=pair):
                plan = self.plan(self.alternating(*pair, lines=24))
                self.assertNotEqual(plan.kind_detected, "transcript",
                                    "an alternating stage marker is still not a speaker")

    def test_prose_that_happens_to_open_with_a_pronoun_in_parentheses_stays_prose(self):
        paragraphs = [
            f"(they said the depot carrier lane {index} would hold, and the tariff on the "
            f"freight side moved twice while the customs desk kept the same note about the "
            f"origin papers that the broker filed for the carrier again) and the review "
            f"closed with the same figure for period {index}."
            for index in range(12)]
        plan = self.plan("\n\n".join(paragraphs))
        self.assertNotEqual(plan.kind_detected, "transcript")

    def test_two_giant_labelled_turns_are_turns_and_are_split_inside(self):
        # The monolithic shape: the whole meeting in two lines, label once at the head.
        first = "(They) " + " ".join(
            f"freight_{index} carrier lane depot tariff customs origin" for index in range(200))
        second = "(You) " + " ".join(
            f"wage_{index} shift roster depot freight carrier tariff margin"
            for index in range(200))
        text = first + "\n" + second + "\n"
        self.assertGreater(len(first), 2500)
        plan = self.plan(text, max_piece_tokens=300)
        self.assertEqual(plan.kind_detected, "transcript",
                         "two labelled turns of any length are a transcript")
        self.assertGreater(len(plan.pieces), 2, "each turn had to be split inside")
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, 300)
        gaps = coverage_gaps(text, plan)
        self.assertEqual(gaps["unaccounted_word_tokens"], 0, gaps)

    def test_one_giant_unlabelled_block_is_not_claimed_as_a_transcript(self):
        text = " ".join(f"freight_{index} carrier lane depot tariff" for index in range(400))
        plan = self.plan(text, max_piece_tokens=300)
        self.assertNotEqual(plan.kind_detected, "transcript")
        self.assertGreater(len(plan.pieces), 1, "it is still cut, just not as turns")




class TestOversizedUnitsAlwaysSplit(unittest.TestCase):
    """No document may fall back because OUR piece was too big.

    Hermes' run lost two documents this way: the structure path applied, then emitted one
    piece over the window, and the fit check sent the file back to LightRAG's default
    chunker - the worst of both, a worse cut and no structure at all.
    """

    WINDOW = 7500
    FLOOR = 150

    def plan(self, text: str, **kwargs):
        kwargs.setdefault("max_piece_tokens", self.WINDOW)
        kwargs.setdefault("min_piece_tokens", self.FLOOR)
        kwargs.setdefault("filename", "Report_2026-04.txt")
        return chunk_document(text, **kwargs)

    def assert_invariant(self, text: str, plan, *, window: int | None = None):
        limit = window or self.WINDOW
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, limit,
                                 f"a piece of {piece.token_estimate} tokens is over the "
                                 f"window; the document would fall back")
        gaps = coverage_gaps(text, plan)
        self.assertEqual(gaps["unaccounted_word_tokens"], 0, gaps)

    def paragraph(self, blocks: int, words: int = 9) -> str:
        return " ".join(f"freight{index} carrier lane depot tariff customs origin papers "
                        f"margin period{index}." for index in range(blocks))

    def test_a_nine_thousand_token_paragraph_is_cut_into_window_sized_pieces(self):
        text = self.paragraph(900)
        self.assertGreater(token_count(text), self.WINDOW)
        plan = self.plan(text)
        self.assertGreaterEqual(len(plan.pieces), 2)
        self.assert_invariant(text, plan)

    def test_the_topic_path_splits_a_unit_it_has_nothing_to_cut_between(self):
        # The bug: with the topic pass on, one unit has no internal topic change, so the
        # segmenter returns it whole and nothing else looks at its size.
        text = self.paragraph(900)
        plan = self.plan(text, topic_embed=lambda texts: [[float(len(t) % 7)] for t in texts],
                         topic_options={"units_per_block": 4})
        self.assertGreaterEqual(len(plan.pieces), 2)
        self.assertGreaterEqual(plan.stats.get("oversized_units_split", 0), 1,
                                "the split is counted, not just performed")
        self.assertNotIn("unit-over-window(split)", plan.warnings,
                         "the topic path is expected to split its own unit; the warning means "
                         "only the last-resort pass saved it")
        self.assert_invariant(text, plan)

    def test_a_twenty_thousand_token_single_line_is_cut_at_sentence_boundaries(self):
        text = self.paragraph(2600)
        plan = self.plan(text)
        self.assert_invariant(text, plan)
        bodies = [piece.body.strip() for piece in plan.pieces]
        for body in bodies[:-1]:
            self.assertTrue(body.endswith((".", "!", "?")),
                            "a sentence boundary was available; use it: " + repr(body[-40:]))

    def test_a_sentenceless_blob_is_hard_split_at_whitespace_never_mid_word(self):
        words = [f"alpha{index}" for index in range(4200)]
        text = " ".join(words)
        plan = self.plan(text)
        self.assert_invariant(text, plan)
        source = set(text.split())
        for piece in plan.pieces:
            self.assertTrue(set(piece.body.split()) <= source,
                            "every piece word is a whole word of the source")

    def test_one_unsplittable_word_is_reported_rather_than_hidden(self):
        monster = "x" * 60000
        text = monster + " " + self.paragraph(60)
        plan = self.plan(text)
        self.assertIn("unit-over-window(one-unsplittable-word)", plan.warnings)
        self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)
        self.assertGreaterEqual(plan.stats.get("units_unsplittable", 0), 1)

    def test_the_pieces_it_does_make_are_balanced_not_one_full_plus_a_stub(self):
        text = self.paragraph(2600)
        plan = self.plan(text)
        sizes = sorted(piece.token_estimate for piece in plan.pieces)
        self.assertGreaterEqual(len(sizes), 3)
        self.assertGreaterEqual(sizes[0], sizes[-1] // 3,
                                f"a piece a third the size of the largest is a stub: {sizes}")
        self.assertLessEqual(sizes[-1], self.WINDOW)

    def test_the_invariant_holds_for_every_kind_this_module_produces(self):
        documents = {
            "headed": "\n".join(f"## Section {index}\n\n" + self.paragraph(60)
                                 for index in range(40)),
            "table": "\n".join(f"| region{index} | 12{index} | 9{index} |"
                                for index in range(1200)),
            "transcript": "\n".join(
                f"(Speaker {index % 3}) " + self.paragraph(30) for index in range(200)),
            "flat": self.paragraph(3000),
        }
        for label, text in documents.items():
            with self.subTest(kind=label):
                plan = self.plan(text)
                self.assert_invariant(text, plan)


class TestMonolithicTwoTurnTranscript(unittest.TestCase):
    """A meeting written as two lines. Nothing repeats, so repetition cannot be the test."""

    WINDOW = 3000

    def turn(self, label: str, marker: str, count: int) -> str:
        return f"({label}) " + " ".join(
            f"{marker}{index} carrier lane depot tariff customs origin papers margin period{index}."
            for index in range(count))

    def plan(self, text: str, **kwargs):
        kwargs.setdefault("max_piece_tokens", self.WINDOW)
        kwargs.setdefault("min_piece_tokens", 150)
        kwargs.setdefault("filename", "Call_2026-04-11.txt")
        return chunk_document(text, **kwargs)

    def test_two_giant_pronoun_turns_are_turns_and_are_split_inside(self):
        text = self.turn("They", "freight", 900) + "\n" + self.turn("You", "wage", 900) + "\n"
        plan = self.plan(text)
        self.assertEqual(plan.kind_detected, "transcript")
        self.assertGreaterEqual(len(plan.pieces), 3, "each turn is over the window on its own")
        for piece in plan.pieces:
            self.assertLessEqual(piece.token_estimate, self.WINDOW)
        sizes = sorted(piece.token_estimate for piece in plan.pieces)
        # A looser band than the flat-path test: a transcript piece must stay inside ONE
        # speaker's turn, so the tail of a turn cannot always be welded to its neighbour.
        self.assertGreaterEqual(sizes[0] * 4, sizes[-1],
                                f"a turn split that leaves a stub retrieves badly: {sizes}")
        self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)

    def test_two_short_labelled_turns_are_still_a_transcript(self):
        # Deliberately under the 2,500-character monolith bar: this document is detected
        # by line shape alone, so it fails if the uniform rule is removed.
        text = self.turn("They", "freight", 24) + "\n" + self.turn("You", "wage", 24) + "\n"
        self.assertLess(len(text.splitlines()[0]), 2500)
        plan = self.plan(text)
        self.assertEqual(plan.kind_detected, "transcript",
                         "every line is a labelled turn; that is the evidence on its own")
        self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)

    def test_a_multi_word_parenthetical_is_not_a_speaker_name(self):
        body_a = ("the carrier lane renewals were signed in the spring, the freight tariff moved "
                  "twice inside one quarter, and the depot foreman kept saying that the numbers "
                  "on the board did not match the numbers in the folder that the broker brought "
                  "around in March")
        body_b = ("the origin papers the broker filed for the third period were never reconciled "
                  "with the depot figures, the customs note was read aloud twice, and nobody in "
                  "the room could say which of the two carriers had actually signed the renewal "
                  "or whether the margin the board saw was the margin anyone had agreed to")
        lines = [f"(Regarding the quarterly numbers) {body_a}",
                 f"(And the customs position) {body_b}"]
        self.assertGreaterEqual(sum(len(line) for line in lines), 600,
                                "long enough that only the label shape can reject this: the "
                                "body-length bar is what would otherwise stop a contents page")
        text = "\n".join(lines) + "\n"
        plan = self.plan(text)
        self.assertNotEqual(plan.kind_detected, "transcript",
                            "two long parentheticals are sentences, not two voices")

    def test_an_unsplittable_word_inside_a_turn_says_which_unit_it_was_in(self):
        monster = "x" * 40000
        text = self.turn("They", "freight", 40) + "\n" + f"(You) {monster} " + self.turn("You", "wage", 10).split(" ", 1)[1] + "\n"
        plan = self.plan(text)
        self.assertEqual(plan.kind_detected, "transcript")
        self.assertIn("unit-over-window(one-unsplittable-word)", plan.warnings,
                      "the warning names the unit that could not be cut, so the row shows "
                      "why a piece is over the window")

    def test_the_uniform_rule_itself_refuses_a_multi_word_label(self):
        # The chunker test above can also be saved by the turn parser, which refuses a
        # multi-word label of its own. This pins the rule that is named in the brief.
        from app.transcript_prep import uniform_label_speaker_set
        lines = ["(Regarding the quarterly numbers) the carrier lane renewals were signed "
                 "in the spring, the freight tariff moved twice inside one quarter, and the "
                 "depot foreman kept saying that the numbers on the board did not match the "
                 "numbers in the folder the broker brought around in March, which is why the "
                 "margin was never signed off by anyone in the room that afternoon.",
                 "(And the customs position) the origin papers filed for the third period "
                 "were never reconciled with the depot figures, the customs note was read "
                 "aloud twice, nobody could say which of the two carriers had actually "
                 "signed the renewal, and the tariff line was carried forward into the next "
                 "quarter's sheet without a single correction being written against it."]
        self.assertGreaterEqual(sum(len(line) for line in lines), 600)
        self.assertEqual(uniform_label_speaker_set(lines), frozenset(),
                         "a four-word parenthetical is a sentence, not a name")

    def test_a_heading_and_a_parenthesised_paragraph_are_not_two_speakers(self):
        text = ("(Introduction)\n(The quarterly report covers both depots and the freight "
                "tariff for the year, the carrier lane renewals signed in the spring, and the "
                "customs position on the origin papers the broker filed for the third period.)\n")
        plan = self.plan(text)
        self.assertNotEqual(plan.kind_detected, "transcript")

    def test_a_contents_page_of_short_paren_labels_is_not_a_transcript(self):
        text = "\n".join(["(Introduction)", "(Scope)", "(Method)", "(Findings)",
                           "(Appendix)"]) + "\n"
        plan = self.plan(text)
        self.assertNotEqual(plan.kind_detected, "transcript")

    def test_one_line_is_never_enough(self):
        plan = self.plan(self.turn("They", "freight", 300) + "\n")
        self.assertNotEqual(plan.kind_detected, "transcript")

    def test_the_split_of_an_over_window_turn_is_balanced_and_whole_word(self):
        from app.structure_chunker import _split_oversize, _Stats
        text = self.turn("They", "freight", 900)
        stats, warnings = _Stats(), []
        parts = _split_oversize(text, 3000, stats, warnings)
        self.assertGreaterEqual(len(parts), 3)
        self.assertEqual(" ".join(parts).split(), text.split(),
                         "the split reorders nothing and loses nothing")
        for part in parts:
            self.assertLessEqual(token_count(part), 3000)
        sizes = sorted(token_count(part) for part in parts)
        self.assertGreaterEqual(sizes[0], sizes[-1] // 3, f"stub pieces: {sizes}")
        self.assertEqual(warnings, [])

    def test_the_speaker_set_rule_itself(self):
        from app.transcript_prep import uniform_label_speaker_set
        lines = [self.turn("They", "freight", 200), self.turn("You", "wage", 200)]
        self.assertEqual(uniform_label_speaker_set(lines), frozenset({"they", "you"}))
        self.assertEqual(uniform_label_speaker_set(["(Introduction)",
                                                    "(Method)"]), frozenset(),
                         "short labels with nothing said in them are a contents page")
        self.assertEqual(uniform_label_speaker_set(["(Cross-talk) the desks talked over each "
                                                    "other for a while at the review of the "
                                                    "carrier lane figure and the customs note "
                                                    "for period three, which nobody wrote down.",
                                                    "(They) the freight figure moved twice."]),
                         frozenset(), "a stage marker is not a speaker")


class TestRotationWithStageMarkers(unittest.TestCase):
    """A real name rotating with a pronoun, with stage directions dropped in between."""

    def transcript(self, lines: int = 120, markers: int = 3) -> str:
        out: list[str] = []
        for index in range(lines):
            label = "Ana" if index % 2 == 0 else "They"
            out.append(f"({label}) the depot carrier lane {index} held while the freight "
                       f"tariff moved twice in the customs review of the origin papers for "
                       f"period {index}, and the {label} note was read again.")
            if markers and index % max(1, lines // markers) == 0:
                out.append("(Cross-talk) both desks spoke at once about the same carrier lane "
                           "figure and the customs note for period three, and neither note was "
                           "written down at the time of the review.")
        return "\n".join(out) + "\n"

    def plan(self, text: str):
        return chunk_document(text, filename="Call_2026-04-11.txt", max_piece_tokens=3000,
                              min_piece_tokens=150)

    def test_the_pronoun_is_rescued_as_the_second_voice(self):
        text = self.transcript()
        plan = self.plan(text)
        self.assertEqual(plan.kind_detected, "transcript")
        bodies = "\n".join(piece.body for piece in plan.pieces)
        self.assertIn("(Ana)", bodies)
        self.assertIn("(They)", bodies)
        self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)

    def test_the_pattern_is_read_after_the_stage_directions_are_dropped(self):
        from app.transcript_prep import speaker_set_from_alternation
        with_markers = self.transcript().splitlines()
        without_markers = [line for line in with_markers if not line.startswith("(Cross-talk)")]
        self.assertEqual(speaker_set_from_alternation(with_markers),
                         frozenset({"ana", "they"}))
        self.assertEqual(speaker_set_from_alternation(without_markers),
                         frozenset({"ana", "they"}))

    def test_a_stage_direction_pair_is_still_never_a_speaker_set(self):
        from app.transcript_prep import speaker_set_from_alternation
        lines = ["(Cross-talk) the desks talked over each other about the carrier lane.",
                 "(Inaudible) the customs note for period three was not recorded."] * 6
        self.assertEqual(speaker_set_from_alternation(lines), frozenset())

    def test_three_rotating_voices_are_not_claimed_by_a_two_voice_rule(self):
        from app.transcript_prep import speaker_set_from_alternation
        lines = []
        for index in range(30):
            label = ("Ana", "They", "We")[index % 3]
            lines.append(f"({label}) the depot carrier lane {index} held while the freight "
                         f"tariff moved twice in the customs review of the origin papers.")
        self.assertEqual(speaker_set_from_alternation(lines), frozenset(),
                         "this rescue is a two-voice rule; three voices need their own")




class TestStubPiecesAreNotPieces(unittest.TestCase):
    """A piece below the merge floor is a mistake, whatever kind of document it came from.

    Hermes' doc01 ended in a 28-token piece: a trailing turn the turn packer could not
    weld because its neighbour was already full, so it stood alone with a header costing
    more than its content.
    """

    WINDOW = 7500
    FLOOR = 150
    STUB = "(Ana) the depot note was read twice, and nobody wrote it down."

    def turn(self, label: str, marker: str, blocks: int) -> str:
        return f"({label}) " + " ".join(
            f"{marker}{i} carrier lane depot tariff customs origin papers margin period{i}."
            for i in range(blocks))

    def plan(self, text: str, **over):
        over.setdefault("max_piece_tokens", self.WINDOW)
        over.setdefault("min_piece_tokens", self.FLOOR)
        over.setdefault("filename", "Call_2026-04-11.txt")
        return chunk_document(text, **over)

    def sizes(self, plan):
        return [piece.token_estimate for piece in plan.pieces]

    def test_a_trailing_stub_turn_welds_into_the_piece_before_it(self):
        text = self.turn("They", "freight", 400) + "\n" + self.STUB + "\n"
        plan = self.plan(text)
        self.assertLess(token_count(self.STUB), self.FLOOR)
        self.assertTrue(all(size >= self.FLOOR for size in self.sizes(plan)),
                        f"a piece under the floor is not a piece: {self.sizes(plan)}")
        self.assertIn("nobody wrote it down", plan.pieces[-1].body)
        self.assertGreaterEqual(plan.stats.get("stub_pieces_merged", 0), 1)
        self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)

    def test_a_label_only_trailing_line_is_welded_the_same_way(self):
        text = self.turn("They", "freight", 400) + "\n(Ana)\n"
        plan = self.plan(text)
        self.assertTrue(all(size >= self.FLOOR for size in self.sizes(plan)),
                        self.sizes(plan))
        self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)

    def test_a_leading_stub_has_no_previous_piece_so_it_welds_forward(self):
        text = self.STUB + "\n" + self.turn("They", "freight", 400) + "\n"
        plan = self.plan(text)
        self.assertTrue(all(size >= self.FLOOR for size in self.sizes(plan)),
                        self.sizes(plan))
        self.assertIn("nobody wrote it down", plan.pieces[0].body,
                      "the first piece is the only neighbour available")

    def test_a_stub_whose_merge_would_break_the_window_is_left_and_says_so(self):
        para = lambda blocks: " ".join(
            f"freight{i} carrier lane depot tariff customs origin papers margin period{i}."
            for i in range(blocks))
        doc = ("## Alpha\n\n" + para(352) + "\n\n## Beta\n\nThe depot note was read twice, "
               "and nobody wrote it down.\n\n## Gamma\n\n" + para(352) + "\n")
        plan = self.plan(doc, filename="Report_2026-04.txt")
        self.assertIn("stub-piece-cannot-merge", plan.warnings)
        self.assertGreaterEqual(plan.stats.get("stub_pieces_left", 0), 1)
        self.assertEqual(coverage_gaps(doc, plan)["unaccounted_word_tokens"], 0,
                         "leaving it alone must never cost text")
        for size in self.sizes(plan):
            self.assertLessEqual(size, self.WINDOW, self.sizes(plan))

    def test_the_weld_prefers_the_previous_neighbour_and_takes_the_next_only_when_it_must(self):
        from app.structure_chunker import _merge_stub_pieces, _Stats
        big = "w " * 3000                     # ~3,000 tokens
        stub = "a short note"
        units = [("section", big, ("Alpha",)), ("paragraph", stub, ()),
                 ("section", big, ("Gamma",))]
        merged = _merge_stub_pieces(units, 150, 7500, _Stats(), [])
        self.assertEqual(len(merged), 2)
        self.assertIn("short note", merged[0][1], "a document reads forward: the previous piece")
        units_first = [("paragraph", stub, ()), ("section", big, ("Alpha",))]
        merged_first = _merge_stub_pieces(units_first, 150, 7500, _Stats(), [])
        self.assertEqual(len(merged_first), 1)
        self.assertIn("short note", merged_first[0][1])
        warnings: list[str] = []
        stuck = _merge_stub_pieces(units, 150, 2900, _Stats(), warnings)
        self.assertEqual(len(stuck), 3, "neither join fits the window, so nothing is merged")
        self.assertIn("stub-piece-cannot-merge", warnings)

    def test_a_document_whose_only_piece_is_small_keeps_it(self):
        text = "(Ana) the depot note was read twice, and nobody wrote it down.\n"
        plan = self.plan(text)
        self.assertEqual(len(plan.pieces), 1)
        self.assertLess(plan.pieces[0].token_estimate, self.FLOOR)

    def test_no_kind_leaves_a_piece_under_the_floor(self):
        para = lambda blocks: " ".join(
            f"freight{i} carrier lane depot tariff customs origin papers margin period{i}."
            for i in range(blocks))
        documents = {
            "transcript trailing stub": (self.turn("They", "freight", 400) + "\n" + self.STUB + "\n",
                                         "Call_2026-04-11.txt"),
            "transcript leading stub": (self.STUB + "\n" + self.turn("They", "freight", 400) + "\n",
                                        "Call_2026-04-11.txt"),
            "headed short tail": ("\n".join(f"## Section {i}\n\n" + para(60) for i in range(20))
                                  + "\n## Tail\n\nThe depot note was read twice.\n",
                                  "Report_2026-04.txt"),
            "flat then a scrap": (para(300) + "\n\nThe depot note was read twice, and nobody "
                                  "wrote it down at the time of the review.\n", "Report_2026-04.txt"),
        }
        for label, (text, filename) in documents.items():
            with self.subTest(shape=label):
                plan = self.plan(text, filename=filename)
                for piece in plan.pieces:
                    self.assertLessEqual(piece.token_estimate, self.WINDOW)
                self.assertEqual(coverage_gaps(text, plan)["unaccounted_word_tokens"], 0)
                if len(plan.pieces) > 1:
                    self.assertTrue(all(piece.token_estimate >= self.FLOOR
                                        for piece in plan.pieces),
                                    f"{label}: {[p.token_estimate for p in plan.pieces]}")


class TestOversizeSplitIsBalanced(unittest.TestCase):
    """A monolithic turn cut in three should come out as three pieces, not two and an orphan."""

    def parts_at(self, text: str, window: int):
        from app.structure_chunker import _split_oversize, _Stats
        stats, warnings = _Stats(), []
        parts = _split_oversize(text, window, stats, warnings)
        return parts, stats, warnings

    def monolith(self, blocks: int) -> str:
        return "(They) " + " ".join(
            f"freight{i} carrier lane depot tariff customs origin papers margin period{i}."
            for i in range(blocks))

    def test_three_pieces_come_out_as_three_equal_pieces(self):
        text = self.monolith(1300)
        parts, stats, warnings = self.parts_at(text, 7500)
        sizes = [token_count(part) for part in parts]
        self.assertEqual(stats.counts["oversized_pieces_rebalanced"], 1)
        self.assertEqual(warnings, [])
        self.assertGreaterEqual(len(sizes), 3)
        self.assertLessEqual(max(sizes), min(sizes) * 1.5,
                             f"pieces should be roughly equal, not a full one and a runt: {sizes}")
        for size in sizes:
            self.assertLessEqual(size, 7500)
        self.assertEqual(" ".join(parts).split(), text.split(),
                         "rebalancing moves boundaries, never words")

    def test_a_split_that_is_already_even_is_left_exactly_as_it_was(self):
        from app.structure_chunker import _rebalance_parts, _Stats
        text = self.monolith(1200)
        words = text.split()
        third = len(words) // 3
        even = [" ".join(words[0:third]), " ".join(words[third:2 * third]),
                " ".join(words[2 * third:])]
        stats = _Stats()
        self.assertEqual(_rebalance_parts(even, 7500, stats), even,
                         "a document with no runt must not move: the benchmark build compares "
                         "piece counts, and an unnecessary change is a regression")
        self.assertEqual(stats.counts["oversized_pieces_rebalanced"], 0)

    def test_the_balance_survives_every_window_that_forces_a_split(self):
        text = self.monolith(1500)
        for window in (7500, 4000, 2500, 1200):
            with self.subTest(window=window):
                parts, _stats, warnings = self.parts_at(text, window)
                self.assertEqual(warnings, [])
                sizes = [token_count(part) for part in parts]
                self.assertTrue(all(size <= window for size in sizes), sizes)
                self.assertEqual(" ".join(parts).split(), text.split())
                if len(sizes) >= 3:
                    self.assertLessEqual(max(sizes), min(sizes) * 2.0, sizes)



class TestTopicPassInTheWorker(unittest.TestCase):
    """The topic pass runs in the worker, and the state row says which kind of cut happened."""

    VOCABULARIES = [["freight", "carrier", "lane"], ["tariff", "duty", "customs"],
                    ["wage", "shift", "roster"]]

    class Embedder:
        def __init__(self, vocabularies):
            self.vocabularies = vocabularies
            self.asks = 0

        def __call__(self, texts):
            self.asks += 1
            return [[1.0 if set(text.lower().split()) & set(topic) else 0.0
                     for topic in self.vocabularies] for text in texts]

    def transcript(self) -> str:
        lines: list[str] = []
        for index, topic in enumerate(self.VOCABULARIES):
            for turn in range(14):
                word = topic[turn % len(topic)]
                other = topic[(turn + 1) % len(topic)]
                lines.append(f"(Speaker {index % 2}) The {word} position sits against "
                             f"{other} for period {turn}, and the {word} margin moved "
                             f"twice while the {other} desk kept the same note.")
        return "\n".join(lines) + "\n"

    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp())
        self.docs = self.root / "source_docs"
        self.docs.mkdir(parents=True)
        self.store = _FakeStore()
        self.post_json = AsyncMock(return_value={"track_id": "trk-1"})
        self.post_file = AsyncMock(return_value={"track_id": "trk-1"})
        patcher = patch.multiple(
            worker.settings,
            source_docs_dir=str(self.docs), uploads_dir=str(self.root / "uploads"),
            transcript_prep_enabled=False, structure_chunking_enabled=True,
            structure_chunking_max_piece_tokens=1200,
            structure_chunking_min_piece_tokens=150,
            structure_chunking_merge_until="floor",
            structure_chunking_header_tokens=20, structure_chunking_headers=True,
            structure_chunking_fit_margin_tokens=64,
            structure_chunking_embed_window_tokens=8192,
            structure_chunking_embed_tokenizer_path="",
            structure_chunking_fail_on_lost_text=False,
            structure_chunking_embed_url="", structure_chunking_embed_model="bge-m3-mlx-fp16",
            structure_chunking_embed_api_key_env="OMLX_API_KEY",
            structure_chunking_embed_batch=4, structure_chunking_embed_timeout_seconds=5.0,
            structure_chunking_topic_threshold=0.30,
            structure_chunking_topic_units_per_window=3,
            structure_chunking_topic_max_windows=100,
            worker_pptx_extract_enabled=True, worker_pdf_extract_enabled=True,
            worker_docx_extract_enabled=True, worker_pdf_text_backend="auto",
            worker_pdf_max_bytes=0, worker_pdf_max_chars=0, worker_docx_max_bytes=0,
            worker_docx_max_chars=0)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _submit(self, name: str, content: str) -> dict[str, Any]:
        path = self.docs / name
        path.write_text(content, encoding="utf-8")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        self.store.docs[str(path)] = {"path": str(path), "sha256": digest, "status": "accepted",
                                      "company": None, "filename": name,
                                      "original_filename": name, "warnings_json": None}
        with patch.multiple(worker,
                            get_document_by_path=self.store.get_document_by_path,
                            get_documents_by_sha256=self.store.get_documents_by_sha256,
                            update_document_state=self.store.update_document_state,
                            upsert_document=self.store.upsert_document,
                            increment_retry=self.store.increment_retry), \
             patch.object(worker.client, "post_json", self.post_json), \
             patch.object(worker.client, "post_file", self.post_file):
            _run_async(worker.submit_file(path, "filesystem", digest))
        return {"path": str(path),
                "json_body": self.post_json.await_args.args[1]
                if self.post_json.await_count else None,
                "meta": self.store.meta(str(path))}

    def test_no_embedder_url_is_recorded_and_the_fallback_is_used(self):
        got = self._submit("Meeting_2026-04-11.txt", self.transcript())
        meta = got["meta"]
        self.assertIn("STRUCTURE_CHUNKING_EMBED_URL",
                      meta["structure_chunking_topic_note"] or "")
        self.assertIsNone(meta.get("structure_chunking_topic_method"),
                          "no topic pass ran: the row must not claim one")
        self.assertEqual(meta["structure_chunking"], "applied")

    def test_the_configured_embedder_cuts_the_flat_transcript_by_topic(self):
        embedder = self.Embedder(self.VOCABULARIES)
        with patch.object(worker, "topic_embedder_for_plan",
                          return_value=(embedder, None)):
            got = self._submit("Meeting_2026-04-11.txt", self.transcript())
        meta = got["meta"]
        self.assertEqual(meta["structure_chunking_topic_method"], "embedder")
        self.assertGreaterEqual(meta["structure_chunking_topic_boundaries"], 2,
                                "three topics, two cuts at minimum")
        self.assertGreaterEqual(embedder.asks, 1, "the call happened in the worker")
        body = got["json_body"]
        delimiter = body["chunking"]["params"]["split_by_character"]
        self.assertGreater(len([part for part in body["text"].split(delimiter) if part]), 2)

    def test_the_merge_mode_is_on_the_row(self):
        got = self._submit("Policy_2026-04.txt",
                           "# Scope\n\n" + "\n\n".join(
                               f"Para {index}: the {word} figure moved in the review."
                               for index, word in enumerate(["freight", "tariff", "wage"]))
                           + "\n")
        self.assertEqual(got["meta"]["structure_chunking_merge_until"], "floor")

    def test_an_embedder_that_raises_still_ingests_the_document(self):
        def broken(texts):
            raise RuntimeError("model server is down")
        with patch.object(worker, "topic_embedder_for_plan", return_value=(broken, None)):
            got = self._submit("Meeting_2026-04-11.txt", self.transcript())
        meta = got["meta"]
        self.assertEqual(meta["structure_chunking"], "applied",
                         "an unreachable model server changes how we cut, not whether we ship")
        self.assertEqual(meta["structure_chunking_topic_method"],
                         "fallback(embedder-unreachable)")



# ── delimiter survival through LightRAG's input sanitiser ──────────────────

class TestDelimiterSurvival(unittest.TestCase):
    def test_no_candidate_can_be_stripped(self):
        # \x1e was the first candidate and the stack2 run proved what it costs: the sanitiser
        # deletes the C0 range, the marker vanishes, ten documents become eight single-segment
        # documents, and the request never says so.
        for candidate in DELIMITER_CANDIDATES:
            self.assertTrue(is_conservative_marker(candidate), candidate)
            self.assertEqual(sanitize_like_lightrag(candidate), candidate.strip(), candidate)
        self.assertNotIn("\x1e", "".join(DELIMITER_CANDIDATES))
        self.assertEqual(DELIMITER_CANDIDATES[0], "\n\n<<<CHUNK-BREAK>>>\n\n",
                         "the default is the documented printable marker")

    def test_an_explicit_delimiter_the_sanitiser_would_delete_is_refused(self):
        plan = chunk_document(TRANSCRIPT, filename="m.txt", budget_tokens=60)
        for marker in ("\x1e", "\n\n\x1e\n\n", "\ufeff", "\u200b", "\x00x"):
            with self.assertRaises(StructureChunkerError, msg=marker) as caught:
                pick_delimiter(plan, marker)
            self.assertIn("sanitiser", str(caught.exception), marker)

    def test_a_whitespace_only_delimiter_is_refused_too(self):
        plan = chunk_document(TRANSCRIPT, filename="m.txt", budget_tokens=60)
        for marker in ("", "   ", "\x0b", "\n\n"):
            with self.assertRaises(StructureChunkerError, msg=marker):
                pick_delimiter(plan, marker)

    def test_an_explicit_delimiter_outside_the_stripped_range_is_obeyed(self):
        plan = chunk_document(TRANSCRIPT, filename="m.txt", budget_tokens=60)
        self.assertEqual(pick_delimiter(plan, "\n\n@@BREAK@@\n\n"), "\n\n@@BREAK@@\n\n")

    def _colliding_plan(self, tail=""):
        # A plan whose bodies literally carry every candidate, marker newlines and all. The
        # chunker cannot normally produce that (it cuts at blank lines), so the picker is
        # tested at its own level: given a document nothing fits, what does it reach for?
        filler = " ".join(f"carrier{index} lane depot tariff customs" for index in range(60))
        pieces = [sc.Piece(body=f"{candidate}{filler}{tail}", header=None, kind="paragraph")
                  for candidate in DELIMITER_CANDIDATES]
        pieces.append(sc.Piece(body=filler + tail, header=None, kind="paragraph"))
        return sc.ChunkPlan(pieces=pieces, kind_detected="plain")

    def test_a_document_that_carries_every_candidate_escalates_by_document_hash(self):
        plan = self._colliding_plan()
        first = pick_delimiter(plan)
        second = pick_delimiter(self._colliding_plan())
        self.assertEqual(first, second, "escalation is deterministic, not random")
        self.assertTrue(is_conservative_marker(first), first)
        self.assertNotIn(first, "\n".join(piece.body for piece in plan.pieces))
        self.assertIn("<<<CHUNK-BREAK-", first)
        self.assertNotIn("\x1e", first, "the escalation never reaches for a control character")

    def test_the_escalated_marker_belongs_to_the_document(self):
        a = pick_delimiter(self._colliding_plan(" alpha"))
        b = pick_delimiter(self._colliding_plan(" beta"))
        self.assertNotEqual(a, b, "two different documents do not share one escalated marker")

    def test_the_escalation_runs_out_of_lengths_before_it_gives_up(self):
        plan = self._colliding_plan()
        marker = pick_delimiter(plan)
        digest = hashlib.sha256("\n".join(p.body for p in plan.pieces).encode()).hexdigest()
        self.assertIn(digest[:8], marker, "the shortest hash marker is tried first")
        # A document that carries even that: the picker keeps going, then refuses to guess.
        clogged = self._colliding_plan()
        clogged.pieces.append(sc.Piece(body="".join(
            f"\n\n<<<CHUNK-BREAK-{digest[:length]}>>>\n\n"
            for length in sc._HASH_DELIMITER_LENGTHS), header=None, kind="paragraph"))
        longer = pick_delimiter(clogged)
        self.assertTrue(is_conservative_marker(longer), longer)

    def test_no_payload_is_marked_at_its_own_edges(self):
        # The whole field is stripped before it is split, so an edge marker loses its newlines
        # and the payload stops carrying the delimiter the request named.
        for name, text in FIXTURES.items():
            plan = chunk_document(text, filename=f"{name}.txt", budget_tokens=120)
            rendered, params = render_for_lightrag(plan)
            marker = params["params"]["split_by_character"]
            self.assertFalse(rendered.startswith(marker), name)
            self.assertFalse(rendered.endswith(marker), name)
            self.assertTrue(sanitize_like_lightrag(rendered).count(marker)
                            == max(0, len(plan.pieces) - 1), name)


class TestChunkTokenSizeFollowsLightRagNotTheEmbedder(unittest.TestCase):
    """doc08 of the stack2 run: cts 2579 against a piece LightRAG counted at 2592.

    ``chunk_token_size`` is LightRAG's own budget, measured with LightRAG's tokenizer family.
    The embedder's count answers a different question (does a piece fit the 8192-token BGE-M3
    window) and using it here sets the budget below what the image measures.
    """

    def _fake(self, name, counter):
        return patch.object(sc, "_TOKENIZER", (name, counter))

    def test_the_size_follows_the_lightrag_counter_when_it_counts_higher(self):
        plan = chunk_document(TRANSCRIPT * 6, filename="m.txt", budget_tokens=60)
        per_char = lambda text: len(text)                     # noqa: E731 - a fake BGE-vs-cl100k
        with self._fake("tiktoken:fake-wide", per_char):
            size, basis = lightrag_chunk_token_size(plan, margin=64)
        self.assertEqual(basis, "tiktoken:fake-wide")
        self.assertEqual(size, max(len(p.render()) for p in plan.pieces) + 64)

    def test_the_size_follows_the_lightrag_counter_when_it_counts_lower(self):
        # A counter that reads LOW is still the right one: LightRAG compares the same numbers,
        # so the budget it enforces is the budget we named. The embedder window stays enforced
        # separately, in the worker, in embedder tokens.
        plan = chunk_document(TRANSCRIPT * 120, filename="m.txt", budget_tokens=60)
        per_hundred = lambda text: max(1, len(text) // 100)   # noqa: E731
        with self._fake("tiktoken:fake-narrow", per_hundred):
            size, basis = lightrag_chunk_token_size(plan, margin=64)
            chars = max(len(p.render()) for p in plan.pieces)
        self.assertEqual(basis, "tiktoken:fake-narrow")
        self.assertLess(size, chars + 64,
                        "the budget follows the counter LightRAG uses even when that counter "
                        "reads far below the embedder's - the embedder window is enforced "
                        "separately, in embedder tokens, in the worker")

    def test_the_basis_is_named_so_a_reader_can_price_the_number(self):
        plan = chunk_document(TRANSCRIPT * 6, filename="m.txt", budget_tokens=60)
        size, basis = lightrag_chunk_token_size(plan, margin=64)
        self.assertEqual(basis, tokenizer_name())
        self.assertIn(basis, ("tiktoken:cl100k_base", "conservative-overcount"),
                      "tiktoken when installed, the documented over-count otherwise - never "
                      "an embedder tokenizer name")
        self.assertNotIn("bge", basis.lower())
        self.assertLessEqual(max(token_count(p.render()) for p in plan.pieces) + 64, size)

    def test_the_size_clears_every_segment_of_the_payload_not_just_every_piece(self):
        plan = chunk_document(TRANSCRIPT * 8, filename="m.txt", budget_tokens=60)
        rendered, params = render_for_lightrag(plan)
        marker = params["params"]["split_by_character"]
        size, _ = chunk_token_size_for_payload(rendered, marker, margin=64)
        for segment in rendered.split(marker):
            if segment.strip():
                self.assertLessEqual(token_count(segment) + 64, size,
                                     "a segment over chunk_token_size is a hard rejection "
                                     "under split_by_character_only, not a re-cut")

    def test_a_payload_segment_bigger_than_any_piece_is_caught_by_the_payload_sizer(self):
        plan = chunk_document(TRANSCRIPT * 8, filename="m.txt", budget_tokens=60)
        piece_size, _ = lightrag_chunk_token_size(plan, margin=64)
        marker = pick_delimiter(plan)
        # A payload that did NOT come from this plan: one segment carries two pieces.
        bloated = (plan.pieces[0].render() + "\n" + plan.pieces[1].render()
                   + marker + plan.pieces[-1].render())
        payload_size, _ = chunk_token_size_for_payload(bloated, marker, margin=64)
        self.assertGreater(payload_size, piece_size,
                           "the payload sizer sees what the plan sizer cannot")

    def test_an_empty_payload_cannot_be_sized(self):
        with self.assertRaises(StructureChunkerError):
            chunk_token_size_for_payload("   ", "\n\n<<<CHUNK-BREAK>>>\n\n", 64)


if __name__ == "__main__":
    unittest.main()

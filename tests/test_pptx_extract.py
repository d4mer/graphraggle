"""Tests for app/pptx_extract.py (packet-27: richer PPTX extraction).

Fixtures are built with python-pptx itself in a temp directory — no binary
.pptx files in git. Nothing here touches the network or the state database.
"""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.pptx_extract import (
    PptxExtractError,
    cap_extracted_text,
    clean_extracted_text,
    extract_pptx_text,
)


def _tmp_dir() -> Path:
    return Path(tempfile.mkdtemp())


def _save(prs: Presentation, name: str = "deck.pptx") -> Path:
    path = _tmp_dir() / name
    prs.save(str(path))
    return path


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, text: str, left: float, top: float, width: float = 3.0, height: float = 1.0):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.text = text
    return box


# ── Slide order, titles, headings ───────────────────────────────────────────

class TestSlideStructure(unittest.TestCase):
    def test_slides_appear_in_order_with_titles(self):
        prs = Presentation()
        first = prs.slides.add_slide(prs.slide_layouts[1])
        first.shapes.title.text = "Q1 revenue"
        first.placeholders[1].text = "Revenue grew 12%."
        second = prs.slides.add_slide(prs.slide_layouts[1])
        second.shapes.title.text = "Outlook"
        second.placeholders[1].text = "Flat in Q2."

        text = extract_pptx_text(_save(prs))
        self.assertLess(text.index("## Slide 1: Q1 revenue"), text.index("## Slide 2: Outlook"))
        self.assertIn("Revenue grew 12%.", text)
        self.assertIn("Flat in Q2.", text)

    def test_title_is_not_repeated_in_the_body(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Q1 revenue"
        slide.placeholders[1].text = "Revenue grew 12%."

        text = extract_pptx_text(_save(prs))
        self.assertEqual(text.count("Q1 revenue"), 1)
        self.assertEqual(text.splitlines()[0], "## Slide 1: Q1 revenue")

    def test_slide_without_title_falls_back_to_untitled(self):
        prs = Presentation()
        _textbox(_blank(prs), "just a box", 1, 1)

        text = extract_pptx_text(_save(prs))
        self.assertIn("## Slide 1: (untitled)", text)

    def test_hidden_slide_is_tagged(self):
        prs = Presentation()
        slide = _blank(prs)
        _textbox(slide, "hidden body", 1, 1)
        slide.element.set("show", "0")

        text = extract_pptx_text(_save(prs))
        self.assertIn("## Slide 1: (untitled) (hidden)", text)

    def test_visible_slide_is_not_tagged(self):
        prs = Presentation()
        _textbox(_blank(prs), "visible body", 1, 1)
        self.assertNotIn("(hidden)", extract_pptx_text(_save(prs)))

    def test_output_is_deterministic(self):
        prs = Presentation()
        slide = _blank(prs)
        _textbox(slide, "right", 5, 2)
        _textbox(slide, "left", 1, 2)
        path = _save(prs)

        self.assertEqual(extract_pptx_text(path), extract_pptx_text(path))


# ── Reading order ───────────────────────────────────────────────────────────

class TestReadingOrder(unittest.TestCase):
    def test_top_to_bottom_then_left_to_right(self):
        prs = Presentation()
        slide = _blank(prs)
        _textbox(slide, "right box", 5, 2)
        _textbox(slide, "left box", 1, 2)
        _textbox(slide, "top box", 1, 0.5)

        text = extract_pptx_text(_save(prs))
        lines = [line for line in text.splitlines() if "box" in line]
        self.assertEqual(lines, ["top box", "left box", "right box"])


# ── Tables ──────────────────────────────────────────────────────────────────

class TestTables(unittest.TestCase):
    def _deck_with_table(self) -> Path:
        prs = Presentation()
        slide = _blank(prs)
        table = slide.shapes.add_table(2, 3, Inches(1), Inches(3), Inches(6), Inches(1.5)).table
        for column, value in enumerate(["Region", "Q1", "QoQ"]):
            table.cell(0, column).text = value
        table.cell(1, 0).text = "APAC"
        table.cell(1, 1).text = "3.4M"
        table.cell(1, 2).text = "+18%\nstrong"
        return _save(prs)

    def test_table_marker_and_one_line_per_row(self):
        text = extract_pptx_text(self._deck_with_table())
        self.assertIn("[Table]", text)
        self.assertIn("| Region | Q1 | QoQ |", text)
        row_line = next(line for line in text.splitlines() if line.startswith("| APAC"))
        self.assertEqual(row_line, "| APAC | 3.4M | +18% strong |")

    def test_cell_newlines_do_not_split_a_row(self):
        text = extract_pptx_text(self._deck_with_table())
        self.assertNotIn("+18%\nstrong", text)
        self.assertEqual(len([line for line in text.splitlines() if "|" in line]), 2)


# ── Grouped shapes ──────────────────────────────────────────────────────────

class TestGroupShapes(unittest.TestCase):
    def test_nested_group_text_is_included(self):
        prs = Presentation()
        slide = _blank(prs)
        outer = slide.shapes.add_group_shape()
        _textbox(outer, "outer group", 1, 1)
        inner = outer.shapes.add_group_shape()
        _textbox(inner, "inner group", 1, 2)

        text = extract_pptx_text(_save(prs))
        self.assertIn("outer group", text)
        self.assertIn("inner group", text)
        self.assertLess(text.index("outer group"), text.index("inner group"))


# ── Speaker notes ───────────────────────────────────────────────────────────

class TestNotes(unittest.TestCase):
    def test_notes_are_included_under_a_marker(self):
        prs = Presentation()
        slide = _blank(prs)
        _textbox(slide, "slide body", 1, 1)
        slide.notes_slide.notes_text_frame.text = "Talk about FX impact.\nAsk for the APAC split."

        text = extract_pptx_text(_save(prs))
        self.assertIn("[Notes]", text)
        notes_block = text.split("[Notes]", 1)[1]
        self.assertIn("Talk about FX impact.", notes_block)
        self.assertIn("Ask for the APAC split.", notes_block)

    def test_slide_without_notes_has_no_notes_marker(self):
        prs = Presentation()
        _textbox(_blank(prs), "no notes here", 1, 1)
        self.assertNotIn("[Notes]", extract_pptx_text(_save(prs)))


# ── Charts (best effort, never fatal) ───────────────────────────────────────

class TestCharts(unittest.TestCase):
    def test_chart_fixture_produces_a_chart_line_without_raising(self):
        prs = Presentation()
        slide = _blank(prs)
        data = CategoryChartData()
        data.categories = ["Jan", "Feb"]
        data.add_series("Revenue", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(2), data)

        text = extract_pptx_text(_save(prs))
        self.assertIn("[Chart]", text)
        self.assertIn("Jan", text)
        self.assertIn("Revenue", text)

    def test_chart_less_deck_has_no_chart_marker(self):
        prs = Presentation()
        _textbox(_blank(prs), "plain slide", 1, 1)
        self.assertNotIn("[Chart]", extract_pptx_text(_save(prs)))


# ── Failure modes and text hygiene ──────────────────────────────────────────

class TestFailuresAndHygiene(unittest.TestCase):
    def test_empty_deck_extracts_to_empty_text(self):
        prs = Presentation()
        self.assertEqual(extract_pptx_text(_save(prs)), "")

    def test_slide_with_only_empty_shapes_extracts_to_empty_text(self):
        prs = Presentation()
        slide = _blank(prs)
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        self.assertEqual(extract_pptx_text(_save(prs)), "")

    def test_garbage_bytes_raise_extract_failed(self):
        path = _tmp_dir() / "garbage.pptx"
        path.write_bytes(b"this is not a zip archive at all")
        with self.assertRaises(PptxExtractError) as caught:
            extract_pptx_text(path)
        self.assertEqual(caught.exception.error_code, "pptx_extract_failed")

    def test_control_characters_are_stripped_and_blank_runs_collapsed(self):
        cleaned = clean_extracted_text("a\x00b\x07\n\n\n\nline\ttail   \n")
        self.assertEqual(cleaned, "ab\n\nline\ttail")

    def test_cap_truncates_with_a_marker(self):
        capped = cap_extracted_text("x" * 500, 100)
        self.assertLessEqual(len(capped), 100 + 60)
        self.assertIn("[...truncated at 100 characters...]", capped)
        self.assertTrue(capped.startswith("x" * 50))

    def test_cap_leaves_short_text_alone(self):
        self.assertEqual(cap_extracted_text("short", 100), "short")
        self.assertEqual(cap_extracted_text("short", 0), "short")


if __name__ == "__main__":
    unittest.main()

"""Structured PPTX text extraction for the ingest worker (packet-27).

LightRAG's built-in extractor (``lightrag/parser/legacy/extractors.py::
_extract_pptx``) concatenates ``shape.text`` of top-level shapes only, so it
drops table cells (graphic frames have no ``.text``), text inside grouped
shapes, speaker notes, and any notion of which slide text came from. The worker
therefore extracts ``.pptx`` itself and submits the result through
``/documents/text`` like a ``.txt`` file.

``extract_pptx_text`` is pure: no network, no database, no settings. Failure
handling is expressed as :class:`PptxExtractError` carrying the worker
``error_code`` so the caller can write the state row without guessing.

python-pptx loads the whole package into memory, so a huge or zip-bomb ``.pptx``
is bounded by the caller's byte limit (``WORKER_PPTX_MAX_BYTES``), not by
anything here.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

# Control characters other than tab/newline/carriage return: they survive into
# the text from pasted content and break line-based downstream chunking.
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_TRAILING_SPACE = re.compile(r"[ \t]+\n")
_BLANK_RUNS = re.compile(r"\n{3,}")

EXTRACT_LIB_MISSING = "pptx_lib_missing"
EXTRACT_FAILED = "pptx_extract_failed"
EXTRACT_EMPTY = "pptx_empty"


class PptxExtractError(Exception):
    """Extraction could not produce usable text.

    ``error_code`` is the value the worker writes to
    ``document_states.error_code``: ``pptx_lib_missing``, ``pptx_extract_failed``
    (corrupt, password-protected, non-ZIP, or over the byte limit) or
    ``pptx_empty``.
    """

    def __init__(self, message: str, *, error_code: str) -> None:
        super().__init__(message)
        self.error_code = error_code


def clean_extracted_text(text: str) -> str:
    """Strip control characters and collapse runs of blank lines."""
    text = _CONTROL_CHARS.sub("", text)
    text = _TRAILING_SPACE.sub("\n", text)
    text = _BLANK_RUNS.sub("\n\n", text)
    return text.strip()


def cap_extracted_text(text: str, max_chars: int | None) -> str:
    """Truncate to ``max_chars`` with a marker line so the loss is visible."""
    if not max_chars or max_chars <= 0 or len(text) <= max_chars:
        return text
    marker = f"\n[...truncated at {max_chars} characters...]"
    keep = max(0, max_chars - len(marker))
    return text[:keep].rstrip() + marker


def extract_pptx_text(path: Path) -> str:
    """Return per-slide structured text for a ``.pptx`` file.

    Layout, in slide order::

        ## Slide N: <title or "(untitled)">
        <body paragraphs, one per line, reading order: top-to-bottom then
        left-to-right by shape position; grouped shapes are recursed into>
        [Table]
        | col1 | col2 | col3 |    (one line per row, cell newlines flattened)
        [Chart] <title>; <categories>; <series names>
        [Notes]
        <speaker notes>

    Raises :class:`PptxExtractError` when the library is missing or the file
    cannot be opened. A deck with no extractable text returns ``""``.
    """
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise PptxExtractError(
            "python-pptx is not installed; cannot extract .pptx content",
            error_code=EXTRACT_LIB_MISSING,
        ) from exc

    try:
        presentation = Presentation(str(path))
    except ImportError as exc:  # pragma: no cover - lxml missing
        raise PptxExtractError(
            "python-pptx is not installed; cannot extract .pptx content",
            error_code=EXTRACT_LIB_MISSING,
        ) from exc
    except Exception as exc:
        # Covers PackageNotFoundError (non-ZIP / corrupt), OLE compound files
        # (password-protected decks), and truncated archives.
        raise PptxExtractError(
            f"python-pptx could not open the file: {type(exc).__name__}: {exc}",
            error_code=EXTRACT_FAILED,
        ) from exc

    blocks: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        block = _slide_block(index, slide)
        if block:
            blocks.append(block)
    return clean_extracted_text("\n\n".join(blocks))


def _slide_block(index: int, slide: Any) -> str:
    title, title_shape_id = _slide_title(slide)
    heading = f"## Slide {index}: {title or '(untitled)'}"
    if _is_hidden(slide):
        heading += " (hidden)"

    lines: list[str] = [heading]
    for shape in _reading_order(slide.shapes):
        lines.extend(_shape_lines(shape, skip_shape_id=title_shape_id))

    notes = _notes_text(slide)
    if notes:
        lines.append("[Notes]")
        lines.append(notes)

    # Drop empty body lines but keep the heading. A slide with neither a title
    # nor any content contributes nothing: emitting a bare heading would both
    # add retrieval noise and stop an effectively empty deck from being reported
    # as pptx_empty.
    body = [line for line in lines[1:] if line.strip()]
    if not body and not title:
        return ""
    return "\n".join([heading, *body]).strip()


def _slide_title(slide: Any) -> tuple[str | None, int | None]:
    """Return the title text and the id of the title shape (if any)."""
    try:
        title_shape = slide.shapes.title
    except Exception:
        return None, None
    if title_shape is None:
        return None, None
    text = _flatten(title_shape.text if getattr(title_shape, "text", None) else "")
    return (text or None), getattr(title_shape, "shape_id", None)


def _is_hidden(slide: Any) -> bool:
    try:
        return slide.element.get("show") == "0"
    except Exception:
        return False


def _reading_order(shapes: Any) -> list[Any]:
    """Top-to-bottom, then left-to-right; deterministic on ties."""

    def key(shape: Any) -> tuple[float, float, int]:
        top = _coerce_number(getattr(shape, "top", None))
        left = _coerce_number(getattr(shape, "left", None))
        shape_id = int(getattr(shape, "shape_id", 0) or 0)
        return (top, left, shape_id)

    try:
        return sorted(shapes, key=key)
    except Exception:
        return list(shapes)


def _coerce_number(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float("inf")


def _shape_lines(shape: Any, *, skip_shape_id: int | None) -> list[str]:
    if skip_shape_id is not None and getattr(shape, "shape_id", None) == skip_shape_id:
        return []

    lines: list[str] = []
    # Grouped shapes: recurse; the group itself carries no text.
    child_shapes = getattr(shape, "shapes", None)
    if getattr(shape, "shape_type", None) is not None and child_shapes is not None and _is_group(shape):
        for child in _reading_order(child_shapes):
            lines.extend(_shape_lines(child, skip_shape_id=None))
        return lines

    if getattr(shape, "has_table", False):
        table_lines = _table_lines(shape)
        if table_lines:
            lines.append("[Table]")
            lines.extend(table_lines)
        return lines

    if getattr(shape, "has_chart", False):
        chart_line = _chart_line(shape)
        if chart_line:
            lines.append(chart_line)
        return lines

    if getattr(shape, "has_text_frame", False):
        for paragraph in _paragraph_lines(getattr(shape, "text_frame", None)):
            lines.append(paragraph)
    return lines


def _is_group(shape: Any) -> bool:
    shape_type = getattr(shape, "shape_type", None)
    name = str(getattr(shape_type, "name", shape_type) or "")
    return "GROUP" in name.upper()


def _paragraph_lines(text_frame: Any) -> list[str]:
    if text_frame is None:
        return []
    lines: list[str] = []
    for paragraph in getattr(text_frame, "paragraphs", []) or []:
        text = _flatten("".join(run.text for run in getattr(paragraph, "runs", []) or []))
        if not text:
            # A paragraph with no runs can still carry text via field/autoshape
            # content; fall back to the frame's own text once.
            continue
        lines.append(text)
    if not lines:
        whole = _flatten(getattr(text_frame, "text", "") or "")
        if whole:
            lines = [line for line in whole.split("\n") if line]
    return lines


def _table_lines(shape: Any) -> list[str]:
    try:
        table = shape.table
    except Exception:
        return []
    lines: list[str] = []
    try:
        rows = list(table.rows)
    except Exception:
        return lines
    for row in rows:
        cells = [_flatten_cell(cell.text) for cell in row.cells]
        if any(cells):
            # Canonical pipe-table row, leading and trailing bar included: app.structure_chunker
            # recognises that shape and repeats the header row when the table splits. A
            # mid-line-only "a | b" reads as prose to the chunker and loses that protection.
            lines.append("| " + " | ".join(cells) + " |")
    return lines


def _chart_line(shape: Any) -> str | None:
    """Best effort: never let a chart fail the whole deck."""
    try:
        chart = shape.chart
    except Exception:
        return None
    parts: list[str] = []
    try:
        if getattr(chart, "has_title", False):
            title = _flatten(chart.chart_title.text_frame.text)
            if title:
                parts.append(title)
    except Exception:
        pass
    try:
        categories = [
            _flatten(str(category))
            for plot in chart.plots
            for category in plot.categories
        ]
        categories = [category for category in categories if category]
        if categories:
            parts.append("; ".join(categories))
    except Exception:
        pass
    try:
        series = [_flatten(str(item.name or "")) for item in chart.series]
        series = [item for item in series if item]
        if series:
            parts.append("; ".join(series))
    except Exception:
        pass
    if not parts:
        return None
    return "[Chart] " + "; ".join(parts)


def _notes_text(slide: Any) -> str:
    try:
        if not getattr(slide, "has_notes_slide", False):
            return ""
        notes_frame = slide.notes_slide.notes_text_frame
    except Exception:
        return ""
    raw = getattr(notes_frame, "text", "") or ""
    raw = _CONTROL_CHARS.sub(" ", raw)
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in raw.splitlines()]
    return "\n".join(line for line in lines if line)


def _flatten(text: str) -> str:
    """Single line: control characters and newlines collapse to spaces."""
    text = _CONTROL_CHARS.sub(" ", text or "")
    return re.sub(r"\s+", " ", text).strip()


def _flatten_cell(text: str) -> str:
    """Table cell: newlines flattened so one row stays one line."""
    return _flatten(text)
